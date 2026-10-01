from __future__ import annotations

import csv
import fcntl
import hashlib
import io
import json
import logging
import os
import re
import secrets
import shutil
import tempfile
import uuid
from collections import Counter
from collections.abc import Iterator, Mapping
from contextlib import contextmanager
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

POLL_STATUSES = frozenset({"draft", "open", "closed"})
QUESTION_TYPES = frozenset({"text", "yes_no", "single_choice", "multiple_choice", "date"})
RESPONSE_MODES = frozenset({"anonymous", "identified"})
POLL_ID_RE = re.compile(r"^[0-9a-f]{32}$")
QUESTION_ID_RE = re.compile(r"^[0-9a-f]{16}$")
MAX_QUESTIONS = 50
MAX_OPTIONS = 50
MAX_ANSWER_TEXT = 5000
MAX_INVITATIONS = 10_000
DEFAULT_RESPONDENT_SESSION_MINUTES = 60
MIN_RESPONDENT_SESSION_MINUTES = 10
MAX_RESPONDENT_SESSION_MINUTES = 24 * 60

log = logging.getLogger("mifp.polls.storage")


class PollError(ValueError):
    pass


class PollNotFound(PollError):
    pass


class PollSessionExpired(PollError):
    pass


def utc_now() -> datetime:
    return datetime.now(UTC)


def _iso_now() -> str:
    return utc_now().isoformat(timespec="seconds")


def _clean_line(value: Any, limit: int) -> str:
    return " ".join(str(value or "").replace("\r", " ").replace("\n", " ").split())[:limit]


def _clean_text(value: Any, limit: int) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()[:limit]


def _polls_root(runtime_config_dir: Path | str) -> Path:
    root = Path(runtime_config_dir) / "polls"
    if root.is_symlink():
        raise PollError("Poll storage cannot be a symbolic link")
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    os.chmod(root, 0o700)
    return root


def _poll_dir(runtime_config_dir: Path | str, poll_id: str, *, create: bool = False) -> Path:
    if not POLL_ID_RE.fullmatch(str(poll_id or "")):
        raise PollNotFound("Poll not found")
    root = _polls_root(runtime_config_dir).resolve()
    candidate = root / poll_id
    if candidate.is_symlink():
        raise PollNotFound("Poll not found")
    if create:
        candidate.mkdir(mode=0o700)
        os.chmod(candidate, 0o700)
    try:
        candidate.resolve().relative_to(root)
    except ValueError as exc:
        raise PollNotFound("Poll not found") from exc
    if not candidate.is_dir():
        raise PollNotFound("Poll not found")
    return candidate


@contextmanager
def _poll_lock(directory: Path, *, shared: bool = False) -> Iterator[None]:
    path = directory / ".lock"
    with path.open("a+b") as handle:
        os.chmod(path, 0o600)
        fcntl.flock(handle.fileno(), fcntl.LOCK_SH if shared else fcntl.LOCK_EX)
        yield


def _atomic_json(path: Path, payload: Any) -> None:
    descriptor, name = tempfile.mkstemp(prefix=f".{path.name}-", suffix=".tmp", dir=path.parent)
    temporary = Path(name)
    try:
        os.fchmod(descriptor, 0o600)
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        os.chmod(path, 0o600)
        directory_fd = os.open(path.parent, os.O_RDONLY)
        try:
            os.fsync(directory_fd)
        finally:
            os.close(directory_fd)
    finally:
        temporary.unlink(missing_ok=True)


def _read_json(path: Path, expected: type) -> Any:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise PollNotFound("Poll not found") from exc
    except (OSError, ValueError, TypeError) as exc:
        raise PollError("Poll metadata is unavailable") from exc
    if not isinstance(payload, expected):
        raise PollError("Poll metadata is malformed")
    return payload


def _parse_moment(value: Any, *, end_of_day: bool = False) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        if len(text) == 10:
            parsed = datetime.combine(
                date.fromisoformat(text), datetime.max.time() if end_of_day else datetime.min.time()
            )
        else:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise PollError("Enter a valid date or date and time") from exc
    return parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else parsed.astimezone(UTC)


def _respondent_session_minutes(value: Any) -> int:
    if value in (None, ""):
        return DEFAULT_RESPONDENT_SESSION_MINUTES
    if isinstance(value, bool):
        raise PollError("Respondent session duration must be a number of minutes")
    try:
        minutes = int(value)
    except (TypeError, ValueError) as exc:
        raise PollError("Respondent session duration must be a number of minutes") from exc
    if minutes < MIN_RESPONDENT_SESSION_MINUTES or minutes > MAX_RESPONDENT_SESSION_MINUTES:
        raise PollError(
            f"Respondent session duration must be between {MIN_RESPONDENT_SESSION_MINUTES} and "
            f"{MAX_RESPONDENT_SESSION_MINUTES} minutes"
        )
    return minutes


def _poll_defaults(payload: dict[str, Any]) -> dict[str, Any]:
    result = dict(payload)
    result["respondent_session_minutes"] = _respondent_session_minutes(
        result.get("respondent_session_minutes", DEFAULT_RESPONDENT_SESSION_MINUTES)
    )
    return result


def _access_expiry(access: Mapping[str, Any]) -> datetime | None:
    value = access.get("expires_at")
    if not value:
        return None
    try:
        return _parse_moment(value)
    except PollError:
        return None


def _assert_access_active(access: Mapping[str, Any]) -> None:
    expires_at = _access_expiry(access)
    if expires_at is None or expires_at <= utc_now():
        raise PollSessionExpired("Your session has expired. Please reopen the invitation link to continue.")


def _question(raw: Mapping[str, Any], *, existing_id: str | None = None) -> dict[str, Any]:
    kind = str(raw.get("type") or "").strip().lower()
    if kind not in QUESTION_TYPES:
        raise PollError("Unsupported question type")
    text = _clean_line(raw.get("question"), 300)
    if not text:
        raise PollError("Every question needs text")
    question_id = str(raw.get("id") or existing_id or secrets.token_hex(8))
    if not QUESTION_ID_RE.fullmatch(question_id):
        question_id = secrets.token_hex(8)
    item: dict[str, Any] = {
        "id": question_id,
        "type": kind,
        "question": text,
        "required": bool(raw.get("required")),
    }
    if kind == "text":
        style = str(raw.get("style") or "single_line")
        item["style"] = style if style in {"single_line", "multi_line"} else "single_line"
    elif kind == "yes_no":
        item["yes_label"] = _clean_line(raw.get("yes_label") or "Yes", 80)
        item["no_label"] = _clean_line(raw.get("no_label") or "No", 80)
    elif kind in {"single_choice", "multiple_choice"}:
        raw_options = raw.get("options")
        if not isinstance(raw_options, list) or not 2 <= len(raw_options) <= MAX_OPTIONS:
            raise PollError("Choice questions need between 2 and 50 options")
        options: list[dict[str, str]] = []
        seen: set[str] = set()
        for raw_option in raw_options:
            if isinstance(raw_option, Mapping):
                label = _clean_line(raw_option.get("label"), 160)
                option_id = str(raw_option.get("id") or secrets.token_hex(8))
            else:
                label = _clean_line(raw_option, 160)
                option_id = secrets.token_hex(8)
            if not label or label.casefold() in seen:
                raise PollError("Choice options must be non-empty and unique")
            if not QUESTION_ID_RE.fullmatch(option_id):
                option_id = secrets.token_hex(8)
            seen.add(label.casefold())
            options.append({"id": option_id, "label": label})
        item["options"] = options
    return item


def validate_poll(raw: Mapping[str, Any], *, existing: Mapping[str, Any] | None = None) -> dict[str, Any]:
    current = dict(existing or {})
    title = _clean_line(raw.get("title"), 180)
    if not title:
        raise PollError("Poll title is required")
    status = str(raw.get("status") or current.get("status") or "draft").lower()
    if status not in POLL_STATUSES:
        raise PollError("Invalid poll status")
    mode = str(raw.get("response_mode") or current.get("response_mode") or "anonymous").lower()
    if mode not in RESPONSE_MODES:
        raise PollError("Invalid response mode")
    deadline = str(raw.get("deadline") or "").strip()
    retention = str(raw.get("retention_until") or "").strip()
    parsed_deadline = _parse_moment(deadline, end_of_day=True) if deadline else None
    if status == "open" and parsed_deadline is None:
        raise PollError("Set a deadline before opening the poll")
    if retention:
        parsed_retention = _parse_moment(retention, end_of_day=True)
        if parsed_deadline and parsed_retention and parsed_retention < parsed_deadline:
            raise PollError("Retention date cannot precede the deadline")
    raw_questions = raw.get("questions")
    if not isinstance(raw_questions, list) or not 1 <= len(raw_questions) <= MAX_QUESTIONS:
        raise PollError("A poll needs between 1 and 50 questions")
    questions = [_question(item) for item in raw_questions if isinstance(item, Mapping)]
    if len(questions) != len(raw_questions):
        raise PollError("Malformed question schema")
    ids = [item["id"] for item in questions]
    if len(ids) != len(set(ids)):
        raise PollError("Question identifiers must be unique")
    now = _iso_now()
    return {
        "id": str(current.get("id") or raw.get("id") or uuid.uuid4().hex),
        "title": title,
        "description": _clean_text(raw.get("description"), 4000),
        "status": status,
        "created_at": str(current.get("created_at") or now),
        "updated_at": now,
        "deadline": deadline,
        "retention_until": retention,
        "response_mode": mode,
        "allow_changes": bool(raw.get("allow_changes", True)),
        "respondent_session_minutes": _respondent_session_minutes(
            raw.get(
                "respondent_session_minutes",
                current.get("respondent_session_minutes", DEFAULT_RESPONDENT_SESSION_MINUTES),
            )
        ),
        "questions": questions,
        "invitation_subject": _clean_line(raw.get("invitation_subject") or title, 160),
        "invitation_message": _clean_text(raw.get("invitation_message"), 10000),
        "invitation_cta": _clean_line(raw.get("invitation_cta") or "Open poll", 60),
        **({"retention_cleaned_at": current["retention_cleaned_at"]} if current.get("retention_cleaned_at") else {}),
    }


def poll_template(name: str) -> dict[str, Any]:
    templates: dict[str, tuple[str, str, list[dict[str, Any]]]] = {
        "blank": (
            "Untitled poll",
            "",
            [{"type": "text", "question": "Your response", "required": False, "style": "multi_line"}],
        ),
        "attendance": (
            "Attendance confirmation",
            "Please confirm whether you will attend.",
            [{"type": "yes_no", "question": "Will you attend?", "required": True}],
        ),
        "availability": (
            "Event availability",
            "Tell us when you are available.",
            [
                {
                    "type": "multiple_choice",
                    "question": "Which dates are suitable?",
                    "required": True,
                    "options": ["Option 1", "Option 2", "Option 3"],
                }
            ],
        ),
        "date": (
            "Choose a date",
            "Choose your preferred date.",
            [{"type": "date", "question": "Which date would you prefer?", "required": True}],
        ),
        "location": (
            "Choose a location",
            "Choose your preferred location.",
            [
                {
                    "type": "single_choice",
                    "question": "Preferred location",
                    "required": True,
                    "options": ["Rome", "Paris", "Geneva"],
                }
            ],
        ),
    }
    title, description, questions = templates.get(name, templates["blank"])
    deadline = (utc_now() + timedelta(days=14)).date().isoformat()
    retention = (utc_now() + timedelta(days=104)).date().isoformat()
    return {
        "title": title,
        "description": description,
        "status": "draft",
        "deadline": deadline,
        "retention_until": retention,
        "response_mode": "anonymous",
        "allow_changes": True,
        "questions": questions,
        "invitation_subject": title,
        "invitation_message": (
            "Hello {{first_name}},\n\n"
            "Please use the secure link below to complete this MIFP poll.\n\n"
            "Thank you."
        ),
        "invitation_cta": "Open poll",
    }


def create_poll(runtime_config_dir: Path | str, template: str = "blank") -> dict[str, Any]:
    poll_id = uuid.uuid4().hex
    directory = _poll_dir(runtime_config_dir, poll_id, create=True)
    payload = validate_poll({**poll_template(template), "id": poll_id})
    with _poll_lock(directory):
        _atomic_json(directory / "poll.json", payload)
        _atomic_json(directory / "invitations.json", {"invitations": []})
        responses = directory / "responses.jsonl"
        responses.touch(mode=0o600, exist_ok=True)
        os.chmod(responses, 0o600)
    return payload


def get_poll(runtime_config_dir: Path | str, poll_id: str) -> dict[str, Any]:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        return _poll_defaults(_read_json(directory / "poll.json", dict))


def save_poll(runtime_config_dir: Path | str, poll_id: str, raw: Mapping[str, Any]) -> dict[str, Any]:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        existing = _poll_defaults(_read_json(directory / "poll.json", dict))
        payload = validate_poll(raw, existing=existing)
        payload["id"] = poll_id
        _atomic_json(directory / "poll.json", payload)
        return payload


def _invitations_unlocked(directory: Path) -> list[dict[str, Any]]:
    payload = _read_json(directory / "invitations.json", dict)
    rows = payload.get("invitations")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise PollError("Invitation metadata is malformed")
    return rows


def _read_responses_unlocked(directory: Path) -> list[dict[str, Any]]:
    path = directory / "responses.jsonl"
    try:
        data = path.read_bytes()
    except FileNotFoundError:
        return []
    if not data:
        return []
    lines = data.splitlines(keepends=True)
    rows: list[dict[str, Any]] = []
    for index, line in enumerate(lines):
        complete = line.endswith((b"\n", b"\r"))
        try:
            value = json.loads(line)
        except (UnicodeDecodeError, ValueError, TypeError) as exc:
            if index == len(lines) - 1 and not complete:
                log.warning("truncated response tail ignored poll_id=%s", directory.name)
                break
            raise PollError("Response storage contains a malformed record") from exc
        if not isinstance(value, dict):
            raise PollError("Response storage contains a malformed record")
        rows.append(value)
    return rows


def _repair_truncated_response_tail(path: Path) -> None:
    """Discard only an incomplete final record left by an interrupted append."""
    try:
        with path.open("r+b") as handle:
            handle.seek(0, os.SEEK_END)
            size = handle.tell()
            if not size:
                return
            handle.seek(-1, os.SEEK_END)
            if handle.read(1) in {b"\n", b"\r"}:
                return
            handle.seek(0)
            data = handle.read()
            boundary = data.rfind(b"\n")
            handle.truncate(boundary + 1 if boundary >= 0 else 0)
            handle.flush()
            os.fsync(handle.fileno())
    except FileNotFoundError:
        return


def current_responses(runtime_config_dir: Path | str, poll_id: str) -> list[dict[str, Any]]:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        rows = _read_responses_unlocked(directory)
    current: dict[str, dict[str, Any]] = {}
    for row in rows:
        invitation_id = str(row.get("invitation_id") or "")
        if not invitation_id:
            continue
        prior = current.get(invitation_id)
        if prior is None or int(row.get("revision") or 0) > int(prior.get("revision") or 0):
            current[invitation_id] = row
    return sorted(current.values(), key=lambda row: str(row.get("submitted_at") or ""), reverse=True)


def list_polls(runtime_config_dir: Path | str) -> list[dict[str, Any]]:
    root = _polls_root(runtime_config_dir)
    cleanup_retention(runtime_config_dir)
    result: list[dict[str, Any]] = []
    for directory in sorted(root.iterdir()):
        if not directory.is_dir() or not POLL_ID_RE.fullmatch(directory.name) or directory.is_symlink():
            continue
        try:
            poll = get_poll(runtime_config_dir, directory.name)
            invitations = invitation_rows(runtime_config_dir, directory.name)
            responses = current_responses(runtime_config_dir, directory.name)
        except PollError:
            continue
        result.append(
            {
                **poll,
                "invitation_count": sum(row.get("status") == "sent" for row in invitations),
                "response_count": len(responses),
            }
        )
    return sorted(result, key=lambda row: str(row.get("updated_at") or ""), reverse=True)


def invitation_rows(runtime_config_dir: Path | str, poll_id: str) -> list[dict[str, Any]]:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        return _invitations_unlocked(directory)


def create_invitation(
    runtime_config_dir: Path | str,
    poll_id: str,
    *,
    first_name: str = "",
    last_name: str = "",
) -> tuple[dict[str, Any], str]:
    directory = _poll_dir(runtime_config_dir, poll_id)
    raw_token = secrets.token_urlsafe(32)
    invitation = {
        "invitation_id": uuid.uuid4().hex,
        "token_hash": hashlib.sha256(raw_token.encode()).hexdigest(),
        "status": "pending",
        "created_at": _iso_now(),
        "sent_at": None,
        "responded_at": None,
        "revoked_at": None,
    }
    with _poll_lock(directory):
        poll = _poll_defaults(_read_json(directory / "poll.json", dict))
        try:
            _assert_poll_accepting(poll)
        except PollError as exc:
            raise PollError("Open the poll with an active deadline before sending invitations") from exc
        if poll.get("response_mode") == "identified":
            invitation["first_name"] = _clean_line(first_name, 120)
            invitation["last_name"] = _clean_line(last_name, 120)
            invitation["display_name"] = _clean_line(f"{first_name} {last_name}", 240)
        rows = _invitations_unlocked(directory)
        if len(rows) >= MAX_INVITATIONS:
            raise PollError("This poll has reached its invitation limit")
        rows.append(invitation)
        _atomic_json(directory / "invitations.json", {"invitations": rows})
    return invitation, raw_token


def set_invitation_status(runtime_config_dir: Path | str, poll_id: str, invitation_id: str, status: str) -> None:
    if status not in {"sent", "revoked", "failed"}:
        raise PollError("Invalid invitation status")
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        rows = _invitations_unlocked(directory)
        for row in rows:
            if secrets.compare_digest(str(row.get("invitation_id") or ""), invitation_id):
                row["status"] = status
                if status == "sent":
                    row["sent_at"] = _iso_now()
                else:
                    row["revoked_at"] = _iso_now()
                    row["token_hash"] = ""
                    for field in ("first_name", "last_name", "display_name"):
                        row.pop(field, None)
                _atomic_json(directory / "invitations.json", {"invitations": rows})
                return
    raise PollError("Invitation not found")


def exchange_token(runtime_config_dir: Path | str, poll_id: str, raw_token: str) -> dict[str, Any]:
    if len(raw_token) < 40 or len(raw_token) > 128:
        raise PollError("Invitation link is invalid or no longer active")
    digest = hashlib.sha256(raw_token.encode()).hexdigest()
    now = utc_now()
    try:
        directory = _poll_dir(runtime_config_dir, poll_id)
        with _poll_lock(directory, shared=True):
            poll = _poll_defaults(_read_json(directory / "poll.json", dict))
            _assert_poll_accepting(poll)
            for row in _invitations_unlocked(directory):
                stored = str(row.get("token_hash") or "")
                if len(stored) == 64 and secrets.compare_digest(stored, digest):
                    if row.get("status") != "sent" or row.get("revoked_at"):
                        break
                    minutes = _respondent_session_minutes(poll.get("respondent_session_minutes"))
                    expires_at = now + timedelta(minutes=minutes)
                    return {
                        "poll_id": poll_id,
                        "invitation_id": str(row["invitation_id"]),
                        "issued_at": now.isoformat(timespec="seconds"),
                        "expires_at": expires_at.isoformat(timespec="seconds"),
                        "session_minutes": minutes,
                    }
    except PollNotFound:
        pass
    except PollError:
        # Do not reveal whether a poll is closed, expired, revoked or merely unknown.
        pass
    raise PollError("Invitation link is invalid or no longer active")


def _assert_poll_accepting(poll: Mapping[str, Any]) -> None:
    if poll.get("status") != "open":
        raise PollError("This poll is not open")
    deadline = _parse_moment(poll.get("deadline"), end_of_day=True)
    if deadline is None or deadline < utc_now():
        raise PollError("This poll deadline has passed")
    retention = _parse_moment(poll.get("retention_until"), end_of_day=True)
    if retention and retention < utc_now():
        raise PollError("This poll is no longer available")


def respondent_poll(runtime_config_dir: Path | str, access: Mapping[str, Any]) -> dict[str, Any]:
    _assert_access_active(access)
    poll_id = str(access.get("poll_id") or "")
    invitation_id = str(access.get("invitation_id") or "")
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        poll = _poll_defaults(_read_json(directory / "poll.json", dict))
        _assert_poll_accepting(poll)
        invitation = next(
            (row for row in _invitations_unlocked(directory) if row.get("invitation_id") == invitation_id), None
        )
        if not invitation or invitation.get("status") != "sent" or invitation.get("revoked_at"):
            raise PollError("Invitation is no longer active")
        current = None
        for row in _read_responses_unlocked(directory):
            if row.get("invitation_id") == invitation_id and (
                current is None or int(row.get("revision") or 0) > int(current.get("revision") or 0)
            ):
                current = row
        if current and not poll.get("allow_changes"):
            return {**poll, "current_answers": current.get("answers", {}), "locked": True}
        return {**poll, "current_answers": current.get("answers", {}) if current else {}, "locked": False}


def _validate_answers(poll: Mapping[str, Any], raw: Any) -> dict[str, Any]:
    if not isinstance(raw, Mapping):
        raise PollError("Answers must be an object")
    questions = poll.get("questions") or []
    allowed_ids = {str(item["id"]) for item in questions}
    if any(str(key) not in allowed_ids for key in raw):
        raise PollError("Answers contain an unknown question")
    answers: dict[str, Any] = {}
    for question in questions:
        qid = str(question["id"])
        kind = question["type"]
        value = raw.get(qid)
        empty = value is None or value == "" or value == []
        if empty:
            if question.get("required"):
                raise PollError(f"Answer required: {question['question']}")
            answers[qid] = [] if kind == "multiple_choice" else ""
            continue
        if kind == "text":
            answer = _clean_text(value, MAX_ANSWER_TEXT)
            if len(str(value)) > MAX_ANSWER_TEXT:
                raise PollError("A text answer is too long")
        elif kind == "yes_no":
            answer = str(value)
            if answer not in {"yes", "no"}:
                raise PollError("Invalid Yes / No answer")
        elif kind in {"single_choice", "multiple_choice"}:
            option_ids = {str(option["id"]) for option in question.get("options", [])}
            submitted = value if isinstance(value, list) else [value]
            if kind == "single_choice" and len(submitted) != 1:
                raise PollError("Select one option")
            if len(submitted) > len(option_ids) or any(str(item) not in option_ids for item in submitted):
                raise PollError("Invalid choice answer")
            answer = [str(item) for item in submitted] if kind == "multiple_choice" else str(submitted[0])
        elif kind == "date":
            try:
                answer = date.fromisoformat(str(value)).isoformat()
            except ValueError as exc:
                raise PollError("Invalid date answer") from exc
        answers[qid] = answer
    return answers


def submit_response(runtime_config_dir: Path | str, access: Mapping[str, Any], raw_answers: Any) -> dict[str, Any]:
    _assert_access_active(access)
    poll_id = str(access.get("poll_id") or "")
    invitation_id = str(access.get("invitation_id") or "")
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        poll = _poll_defaults(_read_json(directory / "poll.json", dict))
        _assert_poll_accepting(poll)
        invitations = _invitations_unlocked(directory)
        invitation = next((row for row in invitations if row.get("invitation_id") == invitation_id), None)
        if not invitation or invitation.get("status") != "sent" or invitation.get("revoked_at"):
            raise PollError("Invitation is no longer active")
        rows = _read_responses_unlocked(directory)
        prior = [row for row in rows if row.get("invitation_id") == invitation_id]
        if prior and not poll.get("allow_changes"):
            raise PollError("This response can no longer be changed")
        answers = _validate_answers(poll, raw_answers)
        submitted_at = _iso_now()
        current_prior = max(prior, key=lambda row: int(row.get("revision") or 0), default=None)
        previous_count = int((current_prior or {}).get("revision_count") or len(prior))
        first_submitted_at = str(
            (current_prior or {}).get("first_submitted_at")
            or (min((str(row.get("submitted_at") or "") for row in prior), default=""))
            or submitted_at
        )
        record = {
            "response_id": uuid.uuid4().hex,
            "invitation_id": invitation_id,
            "revision": max((int(row.get("revision") or 0) for row in prior), default=0) + 1,
            "revision_count": previous_count + 1,
            "first_submitted_at": first_submitted_at,
            "last_modified_at": submitted_at,
            "submitted_at": submitted_at,
            **(
                {"revision_timestamps": [
                    *[str(value) for value in (current_prior or {}).get("revision_timestamps", []) if value],
                    submitted_at,
                ]}
                if (current_prior or {}).get("revision_timestamps")
                else {}
            ),
            "answers": answers,
        }
        path = directory / "responses.jsonl"
        _repair_truncated_response_tail(path)
        descriptor = os.open(path, os.O_WRONLY | os.O_APPEND | os.O_CREAT, 0o600)
        try:
            line = (json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode()
            os.write(descriptor, line)
            os.fsync(descriptor)
        finally:
            os.close(descriptor)
        invitation["responded_at"] = record["submitted_at"]
        _atomic_json(directory / "invitations.json", {"invitations": invitations})
        return record


def _response_groups(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        invitation_id = str(row.get("invitation_id") or "")
        if invitation_id:
            groups.setdefault(invitation_id, []).append(row)
    for records in groups.values():
        records.sort(key=lambda item: (int(item.get("revision") or 0), str(item.get("submitted_at") or "")))
    return groups


def _response_activity(records: list[dict[str, Any]]) -> dict[str, Any]:
    if not records:
        return {"first_submitted_at": "", "last_modified_at": "", "revision_count": 0, "history_retained": 0}
    current = records[-1]
    legacy_compacted = (
        len(records) == 1
        and int(current.get("revision") or 0) > 1
        and not current.get("first_submitted_at")
        and not current.get("revision_timestamps")
    )
    first = "" if legacy_compacted else str(current.get("first_submitted_at") or records[0].get("submitted_at") or "")
    last = str(current.get("last_modified_at") or current.get("submitted_at") or "")
    revision_count = max(int(current.get("revision_count") or 0), int(current.get("revision") or 0), len(records))
    return {
        "first_submitted_at": first,
        "last_modified_at": last,
        "revision_count": revision_count,
        "history_retained": len(records),
    }


def _revision_timestamps(records: list[dict[str, Any]]) -> list[str]:
    if not records:
        return []
    timestamps: list[str] = []
    first_saved = records[0].get("revision_timestamps")
    if isinstance(first_saved, list):
        timestamps.extend(str(value) for value in first_saved if value)
        start = 1
    else:
        start = 0
    for record in records[start:]:
        submitted = str(record.get("submitted_at") or "")
        if submitted:
            timestamps.append(submitted)
    return timestamps


def response_history(runtime_config_dir: Path | str, poll_id: str, response_id: str) -> dict[str, Any]:
    poll = get_poll(runtime_config_dir, poll_id)
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        rows = _read_responses_unlocked(directory)
        invitations = {row["invitation_id"]: row for row in _invitations_unlocked(directory)}
    target = next((row for row in rows if secrets.compare_digest(str(row.get("response_id") or ""), response_id)), None)
    if target is None:
        raise PollNotFound("Response not found")
    invitation_id = str(target.get("invitation_id") or "")
    records = _response_groups(rows).get(invitation_id, [])
    activity = _response_activity(records)
    revisions: list[dict[str, Any]] = []
    previous_answers: Mapping[str, Any] = {}
    for record in records:
        answers = record.get("answers") if isinstance(record.get("answers"), Mapping) else {}
        changed = [
            question["question"]
            for question in poll["questions"]
            if previous_answers and answers.get(question["id"]) != previous_answers.get(question["id"])
        ]
        revisions.append(
            {
                "revision": int(record.get("revision") or 0),
                "submitted_at": str(record.get("submitted_at") or ""),
                "display_answers": {
                    question["id"]: _answer_label(question, answers.get(question["id"]))
                    for question in poll["questions"]
                },
                "changed_questions": changed,
            }
        )
        previous_answers = answers
    identity: dict[str, str] = {}
    if poll.get("response_mode") == "identified":
        invitation = invitations.get(invitation_id, {})
        identity = {
            "first_name": str(invitation.get("first_name") or ""),
            "last_name": str(invitation.get("last_name") or ""),
            "display_name": str(invitation.get("display_name") or ""),
        }
    return {
        **activity,
        **identity,
        "history_compacted": activity["revision_count"] > activity["history_retained"],
        "revision_timestamps": _revision_timestamps(records),
        "revisions": list(reversed(revisions)),
        "questions": [{"id": q["id"], "question": q["question"]} for q in poll["questions"]],
    }


def response_view(runtime_config_dir: Path | str, poll_id: str) -> list[dict[str, Any]]:
    poll = get_poll(runtime_config_dir, poll_id)
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        rows = _read_responses_unlocked(directory)
        invitations = {row["invitation_id"]: row for row in _invitations_unlocked(directory)}
    groups = _response_groups(rows)
    result: list[dict[str, Any]] = []
    for invitation_id, records in groups.items():
        row = records[-1]
        presented = {key: value for key, value in row.items() if key != "invitation_id"}
        presented.update(_response_activity(records))
        presented["display_answers"] = {
            question["id"]: _answer_label(question, row.get("answers", {}).get(question["id"]))
            for question in poll["questions"]
        }
        if poll.get("response_mode") == "identified":
            invitation = invitations.get(invitation_id, {})
            presented["first_name"] = invitation.get("first_name", "")
            presented["last_name"] = invitation.get("last_name", "")
            presented["display_name"] = invitation.get("display_name", "")
        result.append(presented)
    return sorted(
        result,
        key=lambda row: str(row.get("last_modified_at") or row.get("submitted_at") or ""),
        reverse=True,
    )


def poll_analysis(runtime_config_dir: Path | str, poll_id: str) -> dict[str, Any]:
    poll = get_poll(runtime_config_dir, poll_id)
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory, shared=True):
        all_rows = _read_responses_unlocked(directory)
    groups = _response_groups(all_rows)
    responses = [records[-1] for records in groups.values()]
    invitations = invitation_rows(runtime_config_dir, poll_id)
    total = len(responses)
    output: list[dict[str, Any]] = []
    for question in poll["questions"]:
        qid = question["id"]
        values = [row.get("answers", {}).get(qid) for row in responses]
        answered = [value for value in values if value not in (None, "", [])]
        item: dict[str, Any] = {
            "id": qid,
            "question": question["question"],
            "type": question["type"],
            "answered": len(answered),
            "empty": total - len(answered),
        }
        if question["type"] == "yes_no":
            counts = Counter(answered)
            item["options"] = [
                {
                    "label": question.get("yes_label", "Yes"),
                    "value": "yes",
                    "count": counts["yes"],
                    "percentage": round(counts["yes"] * 100 / total, 1) if total else 0,
                },
                {
                    "label": question.get("no_label", "No"),
                    "value": "no",
                    "count": counts["no"],
                    "percentage": round(counts["no"] * 100 / total, 1) if total else 0,
                },
            ]
        elif question["type"] in {"single_choice", "multiple_choice"}:
            flattened = [choice for value in answered for choice in (value if isinstance(value, list) else [value])]
            counts = Counter(flattened)
            item["options"] = [
                {
                    "label": option["label"],
                    "value": option["id"],
                    "count": counts[option["id"]],
                    "percentage": round(counts[option["id"]] * 100 / total, 1) if total else 0,
                }
                for option in question.get("options", [])
            ]
        elif question["type"] == "date":
            counts = Counter(answered)
            item["options"] = [
                {
                    "label": selected,
                    "value": selected,
                    "count": count,
                    "percentage": round(count * 100 / total, 1) if total else 0,
                }
                for selected, count in sorted(counts.items())
            ]
        elif question["type"] == "text":
            item["responses"] = answered
        output.append(item)
    active_invites = sum(row.get("status") == "sent" for row in invitations)
    activity_by_date: dict[str, dict[str, int]] = {}
    modified_responses = 0
    total_revisions = 0
    last_activity = ""
    for records in groups.values():
        info = _response_activity(records)
        revision_count = int(info["revision_count"])
        total_revisions += revision_count
        if revision_count > 1:
            modified_responses += 1
        if str(info["last_modified_at"]) > last_activity:
            last_activity = str(info["last_modified_at"])
        legacy_compacted = (
            len(records) == 1
            and int(records[0].get("revision") or 0) > 1
            and not records[0].get("revision_timestamps")
        )
        for index, submitted in enumerate(_revision_timestamps(records)):
            day = submitted[:10]
            if not day:
                continue
            bucket = activity_by_date.setdefault(day, {"new": 0, "modified": 0})
            if index == 0 and not legacy_compacted:
                bucket["new"] += 1
            else:
                bucket["modified"] += 1
    return {
        "invitations": active_invites,
        "responses": total,
        "response_rate": round(total * 100 / active_invites, 1) if active_invites else 0,
        "modified_responses": modified_responses,
        "total_revisions": total_revisions,
        "last_response_activity": last_activity,
        "activity": [
            {"date": day, "new": counts["new"], "modified": counts["modified"]}
            for day, counts in sorted(activity_by_date.items())
        ],
        "questions": output,
    }


def _answer_label(question: Mapping[str, Any], value: Any) -> str:
    if value in (None, "", []):
        return ""
    if question["type"] == "yes_no":
        return str(question.get("yes_label" if value == "yes" else "no_label") or value)
    if question["type"] in {"single_choice", "multiple_choice"}:
        labels = {option["id"]: option["label"] for option in question.get("options", [])}
        values = value if isinstance(value, list) else [value]
        return "; ".join(labels.get(item, str(item)) for item in values)
    return str(value)


def _spreadsheet_safe(value: Any) -> str:
    text = str(value or "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _export_rows(runtime_config_dir: Path | str, poll_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    poll = get_poll(runtime_config_dir, poll_id)
    rows = response_view(runtime_config_dir, poll_id)
    exported: list[dict[str, Any]] = []
    for row in rows:
        item: dict[str, Any] = {
            "response_id": row["response_id"],
            "first_submitted_at": row.get("first_submitted_at") or row["submitted_at"],
            "last_modified_at": row.get("last_modified_at") or row["submitted_at"],
            "revision_count": int(row.get("revision_count") or row.get("revision") or 1),
        }
        if poll["response_mode"] == "identified":
            item["first_name"] = _spreadsheet_safe(row.get("first_name"))
            item["last_name"] = _spreadsheet_safe(row.get("last_name"))
        for index, question in enumerate(poll["questions"], 1):
            item[f"Q{index} {question['question']}"] = _spreadsheet_safe(
                _answer_label(question, row.get("answers", {}).get(question["id"]))
            )
        exported.append(item)
    return poll, exported


def export_csv(runtime_config_dir: Path | str, poll_id: str) -> bytes:
    poll, rows = _export_rows(runtime_config_dir, poll_id)
    headers = (
        list(rows[0])
        if rows
        else ["response_id", "first_submitted_at", "last_modified_at", "revision_count"]
        + [f"Q{index} {question['question']}" for index, question in enumerate(poll["questions"], 1)]
    )
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def export_xlsx(runtime_config_dir: Path | str, poll_id: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    poll, rows = _export_rows(runtime_config_dir, poll_id)
    analysis = poll_analysis(runtime_config_dir, poll_id)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append(["MIFP Poll report"])
    overview.append(["Title", poll["title"]])
    overview.append(["Status", poll["status"].title()])
    overview.append(["Deadline", poll.get("deadline", "")])
    overview.append(["Invitations", analysis["invitations"]])
    overview.append(["Responses", analysis["responses"]])
    overview.append(["Modified responses", analysis["modified_responses"]])
    overview.append(["Total revisions", analysis["total_revisions"]])
    overview.append(["Last response activity", analysis["last_response_activity"]])
    overview.append(["Response rate", analysis["response_rate"] / 100])
    overview["B10"].number_format = "0.0%"
    responses_sheet = workbook.create_sheet("Responses")
    headers = (
        list(rows[0])
        if rows
        else ["response_id", "first_submitted_at", "last_modified_at", "revision_count"]
        + [f"Q{index} {question['question']}" for index, question in enumerate(poll["questions"], 1)]
    )
    responses_sheet.append(headers)
    for row in rows:
        responses_sheet.append([row.get(header, "") for header in headers])
    responses_sheet.freeze_panes = "A2"
    responses_sheet.auto_filter.ref = responses_sheet.dimensions
    analysis_sheet = workbook.create_sheet("Analysis")
    analysis_sheet.append(["Question", "Option / metric", "Count", "Percentage"])
    for question in analysis["questions"]:
        if question["type"] == "text":
            analysis_sheet.append([question["question"], "Answered", question["answered"], ""])
            analysis_sheet.append([question["question"], "Empty", question["empty"], ""])
        else:
            for option in question.get("options", []):
                analysis_sheet.append(
                    [question["question"], option["label"], option["count"], option["percentage"] / 100]
                )
                analysis_sheet.cell(analysis_sheet.max_row, 4).number_format = "0.0%"
    analysis_sheet.append([])
    activity_header_row = analysis_sheet.max_row + 1
    analysis_sheet.append(["Response activity", "New responses", "Modified responses", ""])
    for activity in analysis["activity"]:
        analysis_sheet.append([activity["date"], activity["new"], activity["modified"], ""])
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="A72B31")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                52, max(12, *(len(str(cell.value or "")) + 2 for cell in column))
            )
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in analysis_sheet[activity_header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="A72B31")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_pdf(runtime_config_dir: Path | str, poll_id: str) -> bytes:
    from xml.sax.saxutils import escape as xml_escape

    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    from mifp_app.services.exporters import MIFP_EXPORT_COLORS

    poll = get_poll(runtime_config_dir, poll_id)
    analysis = poll_analysis(runtime_config_dir, poll_id)
    red = colors.HexColor(f"#{MIFP_EXPORT_COLORS['red']}")
    red_dark = colors.HexColor(f"#{MIFP_EXPORT_COLORS['red_dark']}")
    navy = colors.HexColor(f"#{MIFP_EXPORT_COLORS['navy']}")
    gray_100 = colors.HexColor(f"#{MIFP_EXPORT_COLORS['gray_100']}")
    gray_200 = colors.HexColor(f"#{MIFP_EXPORT_COLORS['gray_200']}")
    gray_500 = colors.HexColor(f"#{MIFP_EXPORT_COLORS['gray_500']}")
    body_color = colors.HexColor("#1F2937")
    display_font = "Times-Bold"
    body_font = "Helvetica"
    body_bold = "Helvetica-Bold"
    body_italic = "Helvetica-Oblique"

    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle(
        "MifpPollBrand",
        parent=styles["Heading1"],
        fontName=display_font,
        fontSize=18,
        leading=20,
        textColor=red,
        spaceAfter=2,
        spaceBefore=0,
    )
    title_style = ParagraphStyle(
        "MifpPollTitle",
        parent=styles["Heading2"],
        fontName=display_font,
        fontSize=14,
        leading=17,
        textColor=navy,
        spaceAfter=3,
    )
    subtitle_style = ParagraphStyle(
        "MifpPollSubtitle",
        parent=styles["Normal"],
        fontName=body_italic,
        fontSize=8.5,
        leading=11,
        textColor=gray_500,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "MifpPollBody",
        parent=styles["BodyText"],
        fontName=body_font,
        fontSize=8.5,
        leading=12,
        textColor=body_color,
    )
    section_style = ParagraphStyle(
        "MifpPollSection",
        parent=styles["Heading2"],
        fontName=display_font,
        fontSize=12,
        leading=15,
        textColor=navy,
        spaceBefore=3 * mm,
        spaceAfter=2 * mm,
        keepWithNext=True,
    )
    question_meta_style = ParagraphStyle(
        "MifpPollQuestionMeta",
        parent=subtitle_style,
        fontName=body_font,
        spaceAfter=2 * mm,
        keepWithNext=True,
    )
    header_style = ParagraphStyle(
        "MifpPollHeader",
        parent=styles["Normal"],
        fontName=body_bold,
        fontSize=7,
        leading=9,
        textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "MifpPollCell",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=7.5,
        leading=10,
        textColor=body_color,
    )
    text_response_style = ParagraphStyle(
        "MifpPollTextResponse",
        parent=body_style,
        backColor=gray_100,
        borderColor=gray_200,
        borderWidth=0.4,
        borderPadding=6,
        spaceAfter=2.5 * mm,
    )

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=22 * mm,
        bottomMargin=18 * mm,
        title=poll["title"],
        author="Mediterranean Institute of Fundamental Physics",
        subject="Poll analysis report",
    )
    exported_at = utc_now().strftime("%Y-%m-%d %H:%M UTC")
    response_label = "response" if analysis["responses"] == 1 else "responses"
    story: list[Any] = [
        Paragraph("MIFP", brand_style),
        Paragraph(xml_escape(str(poll["title"])), title_style),
        Paragraph(
            f"Poll analysis report  ·  Exported: {exported_at}  ·  "
            f"{analysis['responses']} {response_label}",
            subtitle_style,
        ),
        HRFlowable(width="100%", thickness=1.5, color=red, spaceAfter=8, spaceBefore=2),
    ]
    if poll.get("description"):
        story.extend(
            [
                Paragraph(xml_escape(str(poll["description"])), body_style),
                Spacer(1, 4 * mm),
            ]
        )

    summary_headers = ["Status", "Deadline", "Mode", "Invitations", "Responses", "Response rate"]
    summary_values = [
        str(poll["status"]).title(),
        str(poll.get("deadline") or "—"),
        str(poll.get("response_mode") or "anonymous").title(),
        str(analysis["invitations"]),
        str(analysis["responses"]),
        f"{analysis['response_rate']:.1f}%",
    ]
    story.extend(
        [
            Table(
                [
                    [Paragraph(value, header_style) for value in summary_headers],
                    [Paragraph(xml_escape(value), cell_style) for value in summary_values],
                ],
                colWidths=[23 * mm, 32 * mm, 29 * mm, 29 * mm, 27 * mm, 34 * mm],
                repeatRows=1,
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), red),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), body_bold),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, gray_100]),
                        ("GRID", (0, 0), (-1, -1), 0.4, gray_200),
                        ("LINEBELOW", (0, 0), (-1, 0), 1.5, red_dark),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                ),
            ),
            Spacer(1, 3 * mm),
            Paragraph(
                f"Modified responses: <b>{analysis['modified_responses']}</b> &nbsp;&nbsp; "
                f"Total revisions: <b>{analysis['total_revisions']}</b> &nbsp;&nbsp; "
                f"Last activity: <b>{xml_escape(str(analysis['last_response_activity'] or '—'))}</b>",
                body_style,
            ),
            Paragraph("Response activity", section_style),
        ]
    )
    if analysis["activity"]:
        activity_rows = [
            [
                Paragraph("Date (UTC)", header_style),
                Paragraph("New", header_style),
                Paragraph("Modified", header_style),
            ]
        ] + [
            [
                Paragraph(xml_escape(str(item["date"])), cell_style),
                Paragraph(str(item["new"]), cell_style),
                Paragraph(str(item["modified"]), cell_style),
            ]
            for item in analysis["activity"]
        ]
        story.append(
            Table(
                activity_rows,
                repeatRows=1,
                colWidths=[90 * mm, 42 * mm, 42 * mm],
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), red),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.4, gray_200),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                ),
            )
        )
    else:
        story.append(Paragraph("No response activity yet.", body_style))
    story.append(Paragraph("Response analysis", section_style))

    for index, question in enumerate(analysis["questions"], 1):
        story.append(
            Paragraph(
                f"Q{index} · {xml_escape(str(question['question']))}",
                section_style,
            )
        )
        story.append(
            Paragraph(
                f"Answered: {question['answered']}  ·  Empty: {question['empty']}",
                question_meta_style,
            )
        )
        if question["type"] == "text":
            responses = question.get("responses", [])
            if responses:
                for response in responses:
                    story.append(Paragraph(xml_escape(str(response)), text_response_style))
            else:
                story.append(Paragraph("No text responses.", body_style))
        else:
            table_rows = [
                [
                    Paragraph("Option", header_style),
                    Paragraph("Count", header_style),
                    Paragraph("Percentage", header_style),
                ]
            ] + [
                [
                    Paragraph(xml_escape(str(option["label"])), cell_style),
                    Paragraph(str(option["count"]), cell_style),
                    Paragraph(f"{option['percentage']:.1f}%", cell_style),
                ]
                for option in question.get("options", [])
            ]
            story.append(
                Table(
                    table_rows,
                    repeatRows=1,
                    colWidths=[113 * mm, 25 * mm, 36 * mm],
                    style=TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), red),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, 0), body_bold),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, gray_100]),
                            ("GRID", (0, 0), (-1, -1), 0.4, gray_200),
                            ("LINEBELOW", (0, 0), (-1, 0), 1.5, red_dark),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 5),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                            ("TOPPADDING", (0, 0), (-1, -1), 5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ]
                    ),
                )
            )
            options = question.get("options", [])
            if options:
                drawing = Drawing(174 * mm, min(78 * mm, max(32 * mm, len(options) * 11 * mm)))
                chart = HorizontalBarChart()
                chart.x = 45 * mm
                chart.y = 8 * mm
                chart.width = 120 * mm
                chart.height = drawing.height - 14 * mm
                chart.data = [[option["count"] for option in options]]
                chart.categoryAxis.categoryNames = [str(option["label"])[:45] for option in options]
                chart.categoryAxis.labels.fontName = body_font
                chart.categoryAxis.labels.fontSize = 7
                chart.categoryAxis.labels.fillColor = body_color
                chart.categoryAxis.strokeColor = gray_200
                chart.valueAxis.labels.fontName = body_font
                chart.valueAxis.labels.fontSize = 7
                chart.valueAxis.labels.fillColor = gray_500
                chart.valueAxis.strokeColor = gray_200
                chart.valueAxis.valueMin = 0
                chart.valueAxis.valueMax = max(1, *(option["count"] for option in options))
                chart.valueAxis.valueStep = max(1, int(chart.valueAxis.valueMax / 5))
                chart.bars[0].fillColor = red
                drawing.add(chart)
                story.append(Spacer(1, 3 * mm))
                story.append(drawing)
        story.append(Spacer(1, 4 * mm))

    def add_page_footer(canvas, _document):
        page_width, _page_height = A4
        canvas.saveState()
        canvas.setStrokeColor(red)
        canvas.setLineWidth(0.5)
        canvas.line(18 * mm, 13 * mm, page_width - 18 * mm, 13 * mm)
        canvas.setFont(body_font, 7)
        canvas.setFillColor(gray_500)
        canvas.drawString(18 * mm, 8 * mm, "MIFP · Poll analysis")
        canvas.drawRightString(page_width - 18 * mm, 8 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    document.build(story, onFirstPage=add_page_footer, onLaterPages=add_page_footer)
    return output.getvalue()


def clean_poll(runtime_config_dir: Path | str, poll_id: str) -> dict[str, int]:
    """Remove obsolete invitation metadata without discarding response history."""
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        invitations = _invitations_unlocked(directory)
        retained = [
            row
            for row in invitations
            if row.get("status") == "sent" or row.get("responded_at")
        ]
        _atomic_json(directory / "invitations.json", {"invitations": retained})
        return {"revisions_removed": 0, "invitations_removed": len(invitations) - len(retained)}


def anonymize_poll(runtime_config_dir: Path | str, poll_id: str) -> None:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        poll = _poll_defaults(_read_json(directory / "poll.json", dict))
        invitations = _invitations_unlocked(directory)
        for invitation in invitations:
            for key in ("first_name", "last_name", "display_name"):
                invitation.pop(key, None)
        poll["response_mode"] = "anonymous"
        poll["updated_at"] = _iso_now()
        _atomic_json(directory / "invitations.json", {"invitations": invitations})
        _atomic_json(directory / "poll.json", poll)


def reset_poll(runtime_config_dir: Path | str, poll_id: str) -> None:
    directory = _poll_dir(runtime_config_dir, poll_id)
    with _poll_lock(directory):
        _atomic_json(directory / "invitations.json", {"invitations": []})
        response_path = directory / "responses.jsonl"
        descriptor = os.open(response_path, os.O_WRONLY | os.O_TRUNC | os.O_CREAT, 0o600)
        try:
            os.fsync(descriptor)
        finally:
            os.close(descriptor)


def delete_poll(runtime_config_dir: Path | str, poll_id: str) -> None:
    directory = _poll_dir(runtime_config_dir, poll_id)
    root = _polls_root(runtime_config_dir).resolve()
    if directory.resolve().parent != root:
        raise PollNotFound("Poll not found")
    tombstone = root / f".deleted-{uuid.uuid4().hex}"
    with _poll_lock(directory):
        os.replace(directory, tombstone)
    shutil.rmtree(tombstone)


def cleanup_retention(runtime_config_dir: Path | str) -> int:
    root = _polls_root(runtime_config_dir)
    cleaned = 0
    now = utc_now()
    for directory in list(root.iterdir()):
        if not directory.is_dir() or directory.is_symlink() or not POLL_ID_RE.fullmatch(directory.name):
            continue
        try:
            with _poll_lock(directory):
                poll = _poll_defaults(_read_json(directory / "poll.json", dict))
                retention = _parse_moment(poll.get("retention_until"), end_of_day=True)
                if not retention or retention >= now or poll.get("retention_cleaned_at"):
                    continue
                _atomic_json(directory / "invitations.json", {"invitations": []})
                response_path = directory / "responses.jsonl"
                descriptor = os.open(response_path, os.O_WRONLY | os.O_TRUNC | os.O_CREAT, 0o600)
                try:
                    os.fsync(descriptor)
                finally:
                    os.close(descriptor)
                poll["status"] = "closed"
                poll["retention_cleaned_at"] = _iso_now()
                poll["updated_at"] = poll["retention_cleaned_at"]
                _atomic_json(directory / "poll.json", poll)
                cleaned += 1
        except PollError:
            continue
    return cleaned

"""Registration exports reuse the dashboard's institutional document styles."""
from __future__ import annotations

import io
import threading
import math
from pathlib import Path

from PIL import Image, ImageOps, UnidentifiedImageError

from .exporters import rows_to_docx, rows_to_pdf, _sanitize_cell
from .registrations import private_file, proof_filename

# PDFium calls must be serialized even for different documents (upstream API).
_PDF_LOCK = threading.Lock()
MAX_PROOF_BYTES = 10 * 1024 * 1024
MAX_EXPORT_BYTES = 64 * 1024 * 1024
MAX_PROOF_PAGES = 100


def _image_bytes(image):
    if image.width * image.height > 40_000_000:
        raise ValueError('Payment proof image exceeds the pixel limit.')
    image = ImageOps.exif_transpose(image)
    image.thumbnail((1600, 1600))
    rgba = image.convert('RGBA')
    image = Image.new('RGBA', rgba.size, 'white')
    image.alpha_composite(rgba)
    image = image.convert('RGB')
    output = io.BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def payment_proofs(root: Path, event: str, rows: list[dict]) -> list[dict]:
    """Flatten every proof page; missing/invalid proofs stay visible as notices."""
    import pypdfium2 as pdfium

    proofs, total_bytes, page_count, rendered_bytes = [], 0, 0, 0
    def append_image(item, image):
        nonlocal page_count, rendered_bytes
        payload = _image_bytes(image)
        page_count += 1
        rendered_bytes += len(payload)
        if page_count > MAX_PROOF_PAGES or rendered_bytes > MAX_EXPORT_BYTES:
            raise OverflowError('Rendered payment proofs exceed export limits. Apply narrower filters.')
        item['images'].append(payload)
    for row in rows:
        name = proof_filename(row)
        if not name:
            continue
        reference = f"{row['receipt_id']} · {_sanitize_cell(row.get('last_name', ''))[:150]}, {_sanitize_cell(row.get('first_name', ''))[:150]}"
        item = dict(reference=reference, filename=name, images=[], notice='')
        try:
            with private_file(root, 'uploads', event, name) as stream:
                payload = stream.read(MAX_PROOF_BYTES + 1)
            if len(payload) > MAX_PROOF_BYTES:
                raise OverflowError('A payment proof exceeds 10 MiB.')
            total_bytes += len(payload)
            if total_bytes > MAX_EXPORT_BYTES:
                raise OverflowError('Payment proofs exceed 64 MiB. Export a smaller filtered group.')
            if name.endswith('.pdf'):
                with _PDF_LOCK, pdfium.PdfDocument(payload) as document:
                    if not 0 < len(document) <= MAX_PROOF_PAGES - page_count:
                        raise OverflowError('More than 100 proof pages. Export a smaller filtered group.')
                    for index in range(len(document)):
                        page = document[index]
                        try:
                            width, height = page.get_size()
                            if min(width, height) <= 0 or not all(map(math.isfinite, (width, height))):
                                raise ValueError('Invalid PDF page dimensions.')
                            bitmap = page.render(scale=min(2, 1600 / max(width, height)))
                            try:
                                append_image(item, bitmap.to_pil())
                            finally:
                                bitmap.close()
                        finally:
                            page.close()
            else:
                with Image.open(io.BytesIO(payload)) as image:
                    if image.format not in {'PNG', 'JPEG'}:
                        raise ValueError('Unsupported proof image.')
                    append_image(item, image)
        except FileNotFoundError:
            item['notice'] = 'Payment proof file is missing; no copy could be included.'
        except (pdfium.PdfiumError, UnidentifiedImageError, ValueError, Image.DecompressionBombError):
            item['images'] = []
            item['notice'] = 'Payment proof is invalid, encrypted or unsupported; no copy could be included.'
        proofs.append(item)
    return proofs


def registration_document(rows: list[dict], title: str, fmt: str, proofs: list[dict]) -> bytes:
    if len(rows) > 2000:
        raise OverflowError('Document exports support at most 2000 registrations. Apply filters first.')
    # Keep the list legible; the machine-readable CSV/Excel retain all fields.
    columns = ('receipt_id', 'registration_type', 'payment_method', 'admin_status')
    summary = [dict(receipt_id=row['receipt_id'],
                    participant=f"{row.get('last_name', '')}, {row.get('first_name', '')} · {row.get('email', '')}",
                    **{key: row.get(key, '') for key in columns[1:]}) for row in rows]
    if fmt == 'pdf':
        return rows_to_pdf(summary, title, proofs=proofs)
    from docx import Document
    from docx.shared import Cm, RGBColor
    from .exporters import MIFP_EXPORT_COLORS

    document = Document(io.BytesIO(rows_to_docx(summary, title)))
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
    document.styles['Heading 1'].font.name = 'Georgia'
    document.styles['Heading 1'].font.color.rgb = RGBColor.from_string(MIFP_EXPORT_COLORS['red'])
    for proof in proofs:
        document.add_page_break()
        heading = document.add_heading('Proof of payment', level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string(MIFP_EXPORT_COLORS['red'])
        document.add_paragraph(proof['reference'])
        document.add_paragraph(proof['filename'])
        if proof['notice']:
            document.add_paragraph(proof['notice'])
        for index, payload in enumerate(proof['images'], 1):
            if index > 1:
                document.add_page_break()
                document.add_heading('Proof of payment — continued', level=1)
                document.add_paragraph(proof['reference'])
            document.add_paragraph(f"Page {index} of {len(proof['images'])}")
            with Image.open(io.BytesIO(payload)) as image:
                ratio = min(16 / image.width, 19 / image.height)
                document.add_picture(io.BytesIO(payload), width=Cm(image.width * ratio), height=Cm(image.height * ratio))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()

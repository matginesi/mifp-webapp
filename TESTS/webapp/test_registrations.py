from __future__ import annotations

import csv
import importlib.util
import hashlib
import io
import json
import os
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest

from mifp_app.services.registrations import RegistrationAdmin, filtered, private_file, read_records, spreadsheet_safe

ROOT = Path(__file__).resolve().parents[2]
EVENT = 'SYNTHETIC-2027'
ID = 'REG-20261009-AABBCCDD'
FIELDS = ['receipt_id', 'submitted_at', 'first_name', 'last_name', 'email', 'registration_type', 'payment_method', 'proof_file']


@pytest.fixture
def storage(tmp_path):
    private = tmp_path / 'private'
    registrations = private / 'registrations' / EVENT
    uploads = private / 'uploads' / EVENT
    registrations.mkdir(parents=True)
    uploads.mkdir(parents=True)
    (registrations / '.operations.lock').touch()
    (registrations / (ID + '.notifications.lock')).touch()
    row = dict(zip(FIELDS, [ID, '2026-10-09T12:00:00Z', '=1+1', 'Example', 'synthetic@example.invalid', 'student', 'bank', 'proof-' + ID + '.pdf']))
    with (registrations / 'registrations.csv').open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerow(row)
    (uploads / row['proof_file']).write_bytes(b'%PDF-synthetic')
    return private, row


def host_module():
    spec = importlib.util.spec_from_file_location('mifp_registration_tools_test', ROOT / 'deploy/registration-tools.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def host_args(tmp_path, private, row, action='delete'):
    admin = RegistrationAdmin(tmp_path / 'admin')
    request_id = admin.queue(EVENT, [ID], action, 'test-operator', 'participant')
    allowlist = tmp_path / 'allowlist'
    allowlist.write_text('')
    queued = json.loads((admin.directory / 'requests' / (request_id + '.json')).read_text())
    approval = hashlib.sha256(json.dumps(queued, sort_keys=True, separators=(',', ':')).encode()).hexdigest()
    return SimpleNamespace(approve=approval, request=request_id, requests=admin.directory / 'requests', private=private,
                           backups=tmp_path / 'host-backups', confirm=EVENT, allowlist=allowlist,
                           from_address='test@example.invalid', organizer_to='organizer@example.invalid',
                           relay=tmp_path / 'unused-relay')


def test_read_filter_and_metadata_never_change_csv(storage, tmp_path):
    private, row = storage
    path = private / 'registrations' / EVENT / 'registrations.csv'
    original = path.read_bytes()
    admin = RegistrationAdmin(tmp_path / 'metadata')
    admin.update(EVENT, ID, 'confirmed', 'Internal synthetic note', False, 'operator')
    rows = admin.enrich(EVENT, read_records(private, EVENT))
    assert filtered(rows, {'q': 'example', 'status': 'confirmed', 'from': '2026-10-09', 'payment': 'bank'}) == rows
    assert filtered(rows, {'to': '2026-10-08'}) == []
    assert admin.history(EVENT, ID)[0][1] == 'operator'
    assert 'received -> confirmed' in admin.history(EVENT, ID)[0][2]
    assert 'Internal synthetic note' not in admin.history(EVENT, ID)[0][2]
    assert path.read_bytes() == original
    assert spreadsheet_safe(row['first_name']) == "'=1+1"


@pytest.mark.parametrize('payload', [
    'receipt_id,email\nx,a\n',
    ','.join(FIELDS) + '\n' + ','.join([ID] * (len(FIELDS)-1)) + '\n',
    ','.join(FIELDS) + '\n' + ','.join([ID] * len(FIELDS)) + '\n' + ','.join([ID] * len(FIELDS)) + '\n',
])
def test_malformed_or_duplicate_csv_is_refused(storage, payload):
    private, row = storage
    (private / 'registrations' / EVENT / 'registrations.csv').write_text(payload)
    with pytest.raises(ValueError):
        read_records(private, EVENT)


def test_private_files_reject_symlinks_and_traversal(storage, tmp_path):
    private, row = storage
    proof = private / 'uploads' / EVENT / row['proof_file']
    proof.unlink()
    proof.symlink_to('/etc/passwd')
    with pytest.raises(OSError):
        with private_file(private, 'uploads', EVENT, row['proof_file']):
            pass
    with pytest.raises(ValueError):
        with private_file(private, '..', 'passwd'):
            pass
    parent = private / 'registrations' / EVENT
    parent.rename(parent.with_name('real'))
    parent.symlink_to(parent.with_name('real'), target_is_directory=True)
    with pytest.raises(OSError):
        read_records(private, EVENT)


def test_parallel_metadata_writes_and_csv_lock(storage, tmp_path):
    private, row = storage
    admin = RegistrationAdmin(tmp_path / 'metadata')
    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(lambda n: admin.update(EVENT, ID, 'received', f'note {n}', False, 'operator'), range(12)))
    assert len(admin.history(EVENT, ID)) == 12
    # A separate process must block on the same flock used by PHP.
    import fcntl
    with (private / 'registrations' / EVENT / 'registrations.csv').open('rb') as stream:
        fcntl.flock(stream, fcntl.LOCK_EX)
        code = 'import fcntl,sys; f=open(sys.argv[1]); fcntl.flock(f,fcntl.LOCK_SH); print("read")'
        process = subprocess.Popen(['python3', '-c', code, stream.name], stdout=subprocess.PIPE)
        try:
            with pytest.raises(subprocess.TimeoutExpired):
                process.communicate(timeout=0.2)
            fcntl.flock(stream, fcntl.LOCK_UN)
            assert process.communicate(timeout=5)[0] == b'read\n'
        finally:
            if process.poll() is None:
                process.kill()
                process.wait()


def test_delete_requires_revocation_backup_and_can_restore(storage, tmp_path):
    private, row = storage
    args = host_args(tmp_path, private, row)
    host = host_module()
    args.allowlist.write_text(EVENT + '/regform\n')
    with pytest.raises(ValueError, match='Revoke'):
        host.apply(args)
    assert read_records(private, EVENT)
    args.allowlist.write_text('')
    host.apply(args)
    assert read_records(private, EVENT) == []
    assert not (private / 'uploads' / EVENT / row['proof_file']).exists()
    assert (args.backups / (args.request + '.tar')).stat().st_mode & 0o777 == 0o600
    host.apply(args)  # completed requests are idempotent
    host.restore(args)
    assert read_records(private, EVENT)[0] == row
    assert (private / 'uploads' / EVENT / row['proof_file']).read_bytes() == b'%PDF-synthetic'
    with pytest.raises(ValueError, match='duplicate'):
        host.restore(args)


def test_backup_failure_prevents_deletion(storage, tmp_path, monkeypatch):
    private, row = storage
    args = host_args(tmp_path, private, row)
    host = host_module()
    monkeypatch.setattr(host, 'backup', lambda *args: (_ for _ in ()).throw(OSError('synthetic disk full')))
    with pytest.raises(OSError):
        host.apply(args)
    assert read_records(private, EVENT) == [row]


@pytest.mark.parametrize('returncode, expected', [(0, 'accepted'), (1, 'unknown')])
def test_retry_is_explicit_and_never_replays_ambiguous_attempt(storage, tmp_path, monkeypatch, returncode, expected):
    private, row = storage
    args = host_args(tmp_path, private, row, 'retry')
    state = private / 'registrations' / EVENT / (ID + '.notifications.json')
    state.write_text(json.dumps({'participant': {'status': 'failed'}}))
    host = host_module()
    calls = []
    def send(command, **kwargs):
        calls.append(command)
        return SimpleNamespace(returncode=returncode)
    monkeypatch.setattr(host.subprocess, 'run', send)
    host.apply(args)
    host.apply(args)
    assert len(calls) == 1
    assert json.loads(state.read_text())['participant']['status'] == expected
    other = RegistrationAdmin(tmp_path / 'admin').queue(EVENT, [ID], 'retry', 'operator', 'participant')
    args.request = other
    queued = json.loads((args.requests / (other + '.json')).read_text())
    args.approve = hashlib.sha256(json.dumps(queued, sort_keys=True, separators=(',', ':')).encode()).hexdigest()
    with pytest.raises(ValueError, match='Retry blocked'):
        host.apply(args)
    assert len(calls) == 1


@pytest.fixture
def registration_client(tmp_path, storage):
    from mifp_app import create_app
    from mifp_app.db.manage import init_database
    from mifp_app.db.connection import connect
    app = create_app()
    database = tmp_path / 'app.db'
    init_database(database)
    app.config.update(TESTING=True, DATABASE_PATH=database, CONFERENCES_DIR=tmp_path / 'conferences',
                      REGISTRATIONS_PRIVATE_ROOT=str(storage[0]), CSRF_ENABLED=True)
    from werkzeug.security import generate_password_hash
    app.config['ADMIN_PASSWORD_HASH'] = generate_password_hash('synthetic-operation-password', method='pbkdf2:sha256:1000')
    with connect(database) as conn:
        conn.execute("INSERT INTO conference_sites(slug,title,public_path) VALUES(?,?,?)", ('synthetic', 'Synthetic event', EVENT))
        conn.commit()
    client = app.test_client()
    return app, client


def authenticate(client):
    with client.session_transaction() as session:
        session['admin_logged_in'] = True
        session['admin_username'] = 'operator'
        session['_csrf_token'] = 'synthetic-csrf'


def test_dashboard_auth_filters_export_detail_and_csrf(registration_client):
    app, client = registration_client
    base = '/dashboard/conference-sites/1'
    assert client.get(base + '/registrations').status_code in {302, 401}
    authenticate(client)
    result = client.get(base + '/registrations?q=example')
    assert result.status_code == 200, result.data.decode()
    assert b'1 matching / 1 saved' in result.data
    assert b'no-store' in result.headers['Cache-Control'].encode()
    assert client.get(base + '/registrations/' + ID).status_code == 200
    assert client.get(base + '/registrations/' + ID + '/proof').data == b'%PDF-synthetic'
    exported = client.post(base + '/registrations-export/csv', data={'_csrf_token': 'synthetic-csrf', 'password': 'synthetic-operation-password'})
    assert b"'=1+1" in exported.data
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(client.post(base + '/registrations-export/xlsx', data={'_csrf_token': 'synthetic-csrf', 'password': 'synthetic-operation-password'}).data))
    assert workbook.active.cell(5, 3).data_type != 'f'
    assert client.post(base + '/registrations/' + ID, data={'status': 'confirmed'}).status_code == 400
    authenticate(client)  # Failed CSRF requests rotate the session token.
    assert client.post(base + '/registrations/' + ID, data={'status': 'confirmed', '_csrf_token': 'synthetic-csrf'}).status_code == 302
    assert client.get(base + '/registrations?status=cancelled').status_code == 200
    assert b'0 matching' in client.get(base + '/registrations?status=cancelled').data


def test_dashboard_requests_require_confirmation_and_known_notifications(registration_client):
    app, client = registration_client
    authenticate(client)
    endpoint = '/dashboard/conference-sites/1/registrations-request'
    data = {'_csrf_token': 'synthetic-csrf', 'receipt': ID, 'action': 'delete', 'confirm': 'wrong', 'password': 'synthetic-operation-password'}
    assert client.post(endpoint, data=data).status_code == 400
    data.update(confirm=EVENT, action='retry', channel='participant')
    assert client.post(endpoint, data=data).status_code == 409
    data.update(action='delete')
    assert client.post(endpoint, data=data).status_code == 302
    assert read_records(Path(app.config['REGISTRATIONS_PRIVATE_ROOT']), EVENT)
    app.config['REGISTRATIONS_PRIVATE_ROOT'] = ''
    assert client.get('/dashboard/conference-sites/1/registrations').status_code == 200


def test_host_cleanup_rejects_symlink_in_privileged_source(storage, tmp_path):
    private, row = storage
    args = host_args(tmp_path, private, row)
    host = host_module()
    target = private / 'uploads' / EVENT / row['proof_file']
    target.unlink()
    target.symlink_to('/etc/passwd')
    with pytest.raises(OSError):
        host.apply(args)
    assert read_records(private, EVENT) == [row]


def test_registration_access_is_narrow_and_requires_revocation(storage, tmp_path):
    private, row = storage
    host = host_module()
    args = host_args(tmp_path, private, row)
    args.request = EVENT
    args.php_uid, args.php_gid = os.getuid(), os.getgid()
    unrelated = private / 'sessions'
    unrelated.mkdir(mode=0o700)
    session_file = unrelated / 'synthetic-session'
    session_file.write_text('synthetic')
    session_file.chmod(0o600)
    args.allowlist.write_text(EVENT + '/regform\n')
    with pytest.raises(ValueError):
        host.grant_read(args)
    args.allowlist.write_text('')
    host.grant_read(args)
    assert (private / 'registrations' / EVENT / 'registrations.csv').stat().st_mode & 0o777 == 0o640
    assert session_file.stat().st_mode & 0o777 == 0o600
    assert unrelated.stat().st_mode & 0o777 == 0o700


def test_operation_backup_retention_preserves_incomplete_recovery(tmp_path):
    host = host_module()
    for index in range(4):
        identifier = f'{index:032x}'
        archive = tmp_path / (identifier + '.tar')
        archive.write_bytes(b'synthetic')
        os.utime(archive, (index + 1, index + 1))
        (tmp_path / (identifier + '.result.json')).write_text(json.dumps(dict(
            action='delete', status='started' if index == 0 else 'completed')))
    host.prune_backups(tmp_path, 2)
    assert (tmp_path / ('0' * 32 + '.tar')).exists()
    assert not (tmp_path / (f'{1:032x}' + '.tar')).exists()
    assert len(list(tmp_path.glob('*.tar'))) == 3


def test_legacy_unknown_notification_requires_explicit_host_attestation(storage, tmp_path):
    private, row = storage
    host = host_module()
    args = host_args(tmp_path, private, row)
    args.request, args.resolve_receipt, args.channel, args.not_accepted = EVENT, ID, 'participant', False
    with pytest.raises(ValueError):
        host.resolve_notification(args)
    args.not_accepted = True
    host.resolve_notification(args)
    state = private / 'registrations' / EVENT / (ID + '.notifications.json')
    assert json.loads(state.read_text())['participant']['status'] == 'failed'
    state.write_text(json.dumps({'participant': {'status': 'accepted'}}))
    with pytest.raises(ValueError, match='cannot be reset'):
        host.resolve_notification(args)


def test_cleanup_detects_nested_php_authorizations(storage, tmp_path):
    private, row = storage
    args = host_args(tmp_path, private, row)
    args.allowlist.write_text('nested/' + EVENT + '/regform\n')
    with pytest.raises(ValueError, match='Revoke'):
        host_module().apply(args)


def test_ambiguous_conference_storage_mapping_is_blocked(registration_client):
    app, client = registration_client
    from mifp_app.db.connection import connect
    with connect(app.config['DATABASE_PATH']) as conn:
        conn.execute('INSERT INTO conference_sites(slug,title,public_path) VALUES(?,?,?)', ('other', 'Other', 'nested/' + EVENT))
        conn.commit()
    authenticate(client)
    assert client.get('/dashboard/conference-sites/1/registrations').status_code == 409


def test_host_approval_is_bound_to_reviewed_request_contents(storage, tmp_path):
    private, row = storage
    args = host_args(tmp_path, private, row)
    request_path = args.requests / (args.request + '.json')
    value = json.loads(request_path.read_text())
    value['action'] = 'retry'
    request_path.write_text(json.dumps(value))
    with pytest.raises(ValueError, match='Review and supply'):
        host_module().apply(args)
    assert read_records(private, EVENT) == [row]


def test_pagination_and_ordering_preserve_filters(registration_client):
    app, client = registration_client
    private = Path(app.config['REGISTRATIONS_PRIVATE_ROOT'])
    path = private / 'registrations' / EVENT / 'registrations.csv'
    with path.open('a', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        for index in range(26):
            writer.writerow(dict(receipt_id=f'REG-20261009-{index:08X}', submitted_at='2026-10-09',
                                 first_name='Synthetic', last_name=f'Example {index:02}', email='synthetic@example.invalid',
                                 registration_type='student', payment_method='bank', proof_file=''))
    authenticate(client)
    result = client.get('/dashboard/conference-sites/1/registrations?q=example&sort=last_name&order=asc&page=2')
    assert result.status_code == 200
    assert b'Page 2 / 2' in result.data
    assert b'Example 25' in result.data
    assert b'order=asc' in result.data


def test_ordinary_status_normalizes_legacy_mail_provider(tmp_path):
    from TESTS.webapp.test_deploy_state_machine import _env, _run
    env, home = _env(tmp_path)
    config = Path(env['MIFP_CONFIG_DIR']) / 'config.env'
    with config.open('a') as stream:
        stream.write('MAIL_PROVIDER=enabled\nSMTP_HOST=synthetic.example.invalid\n')
    result = _run(env, 'status')
    assert 'Notifications:' in result.stdout
    assert 'Notifications:      mail disabled' not in result.stdout


def test_corrupted_cleanup_backup_cannot_restore(storage, tmp_path):
    private, row = storage
    host = host_module()
    args = host_args(tmp_path, private, row)
    host.apply(args)
    archive = args.backups / (args.request + '.tar')
    with archive.open('ab') as stream:
        stream.write(b'synthetic-corruption')
    with pytest.raises(ValueError, match='integrity'):
        host.restore(args)
    assert read_records(private, EVENT) == []


@pytest.mark.parametrize('writable', [False, True])
def test_host_initializes_removed_event_storage_without_dummy_records(storage, tmp_path, writable):
    import shutil
    private, row = storage
    host = host_module()
    args = host_args(tmp_path, private, row)
    args.request = EVENT
    args.php_uid, args.php_gid = os.getuid(), os.getgid()
    args.writable = writable
    args.public_root = tmp_path/'published'
    source = args.public_root/EVENT/'regform/src'
    source.mkdir(parents=True)
    (source/'bootstrap.php').write_text('function registration_file_mode()')
    (source.parent/'index.php').write_text('$registrationOperationLock .operations.lock')
    args.allowlist.write_text('')
    for tree in ('registrations', 'uploads'):
        shutil.rmtree(private/tree/EVENT)
    host.grant_read(args)
    host.grant_read(args)
    assert read_records(private, EVENT) == []
    assert not (private/'registrations'/EVENT/'registrations.csv').exists()
    for tree in ('registrations', 'uploads'):
        directory = private/tree/EVENT
        assert directory.stat().st_mode & 0o7777 == (0o2770 if writable else 0o750)
    assert (private/'registrations'/EVENT/'.operations.lock').stat().st_mode & 0o777 == (0o660 if writable else 0o640)

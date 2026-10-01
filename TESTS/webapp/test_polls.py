from __future__ import annotations

import hashlib
import io
import re
import stat
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash


@pytest.fixture
def poll_app(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("TESTING", "1")
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "mifp.db"))
    monkeypatch.setenv("ASSETS_DIR", str(tmp_path / "assets"))
    monkeypatch.setenv("EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setenv("CONFERENCES_DIR", str(tmp_path / "conferences"))
    monkeypatch.setenv("RUNTIME_CONFIG_DIR", str(tmp_path / "config"))
    monkeypatch.setenv("LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.setenv("SECRET_KEY", "poll-test-secret-key-that-is-long-enough")
    from mifp_app import create_app
    from mifp_app.db.manage import init_database

    init_database(tmp_path / "mifp.db")
    app = create_app()
    app.config.update(
        TESTING=True,
        WTF_CSRF_ENABLED=False,
        DATABASE_PATH=tmp_path / "mifp.db",
        ASSETS_DIR=tmp_path / "assets",
        EXPORT_DIR=tmp_path / "exports",
        CONFERENCES_DIR=tmp_path / "conferences",
        RUNTIME_CONFIG_DIR=tmp_path / "config",
        LOG_DIR=tmp_path / "logs",
        ADMIN_USERNAME="admin",
        ADMIN_PASSWORD_HASH=generate_password_hash("secret123"),
        MAIL_PROVIDER="console",
        MAIL_TO="operator@example.org",
        MAIL_FROM="alerts@mifp.eu",
        ENV="test",
    )
    for key in (
        "ASSETS_DIR",
        "EXPORT_DIR",
        "CONFERENCES_DIR",
        "RUNTIME_CONFIG_DIR",
        "LOG_DIR",
    ):
        Path(app.config[key]).mkdir(parents=True, exist_ok=True)
    return app


def logged_in_client(app):
    client = app.test_client()
    with client.session_transaction() as session:
        session["admin_logged_in"] = True
        session["admin_username"] = "admin"
        session["admin_login_at"] = time.time()
    return client


def open_poll(runtime: Path, template: str = "attendance"):
    from mifp_app.services.polls import create_poll, save_poll

    poll = create_poll(runtime, template)
    poll["status"] = "open"
    return save_poll(runtime, poll["id"], poll)


def create_sent_invitation(runtime: Path, poll_id: str, **kwargs):
    from mifp_app.services.polls import create_invitation, set_invitation_status

    invitation, token = create_invitation(runtime, poll_id, **kwargs)
    set_invitation_status(runtime, poll_id, invitation["invitation_id"], "sent")
    return invitation, token


def test_poll_storage_is_separate_restrictive_and_token_hash_only(
    tmp_path: Path,
) -> None:
    from mifp_app.services.polls import create_invitation, invitation_rows

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    _invitation, token = create_invitation(
        runtime, poll["id"], first_name="Ada", last_name="Lovelace"
    )
    directory = runtime / "polls" / poll["id"]

    assert stat.S_IMODE(directory.stat().st_mode) == 0o700
    assert stat.S_IMODE((directory / "poll.json").stat().st_mode) == 0o600
    assert stat.S_IMODE((directory / "invitations.json").stat().st_mode) == 0o600
    assert stat.S_IMODE((directory / "responses.jsonl").stat().st_mode) == 0o600
    assert not list(tmp_path.glob("**/*.db"))
    persisted = (directory / "invitations.json").read_text(encoding="utf-8")
    assert token not in persisted
    assert hashlib.sha256(token.encode()).hexdigest() in persisted
    assert "Ada" not in persisted  # anonymous is the default
    assert "email" not in persisted.lower()
    assert invitation_rows(runtime, poll["id"])[0]["status"] == "pending"


def test_all_question_types_revision_and_analysis(tmp_path: Path) -> None:
    from mifp_app.services.polls import (
        create_poll,
        exchange_token,
        poll_analysis,
        save_poll,
        submit_response,
    )

    runtime = tmp_path / "config"
    poll = create_poll(runtime)
    poll.update(
        status="open",
        questions=[
            {
                "id": "1" * 16,
                "type": "text",
                "question": "Comment",
                "required": False,
                "style": "multi_line",
            },
            {
                "id": "2" * 16,
                "type": "yes_no",
                "question": "Attend?",
                "required": True,
                "yes_label": "Available",
                "no_label": "Not available",
            },
            {
                "id": "3" * 16,
                "type": "single_choice",
                "question": "Location",
                "required": True,
                "options": [
                    {"id": "a" * 16, "label": "Rome"},
                    {"id": "b" * 16, "label": "Paris"},
                ],
            },
            {
                "id": "4" * 16,
                "type": "multiple_choice",
                "question": "Dates",
                "required": True,
                "options": [
                    {"id": "c" * 16, "label": "Monday"},
                    {"id": "d" * 16, "label": "Tuesday"},
                ],
            },
            {
                "id": "5" * 16,
                "type": "date",
                "question": "Preferred date",
                "required": True,
            },
        ],
    )
    poll = save_poll(runtime, poll["id"], poll)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    access = exchange_token(runtime, poll["id"], token)
    answers = {
        "1" * 16: "=private note",
        "2" * 16: "yes",
        "3" * 16: "a" * 16,
        "4" * 16: ["c" * 16, "d" * 16],
        "5" * 16: "2027-01-20",
    }
    first = submit_response(runtime, access, answers)
    answers["2" * 16] = "no"
    second = submit_response(runtime, access, answers)
    analysis = poll_analysis(runtime, poll["id"])

    assert (first["revision"], second["revision"]) == (1, 2)
    assert analysis["responses"] == 1
    assert analysis["questions"][1]["options"][1]["count"] == 1
    assert [item["count"] for item in analysis["questions"][3]["options"]] == [1, 1]
    assert (
        sum(item["percentage"] for item in analysis["questions"][3]["options"]) == 200
    )
    lines = (
        (runtime / "polls" / poll["id"] / "responses.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    )
    assert len(lines) == 2


def test_identified_mode_keeps_names_but_never_email(tmp_path: Path) -> None:
    from mifp_app.services.polls import (
        exchange_token,
        response_view,
        save_poll,
        submit_response,
    )

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    poll["response_mode"] = "identified"
    poll = save_poll(runtime, poll["id"], poll)
    _invitation, token = create_sent_invitation(
        runtime, poll["id"], first_name="Grace", last_name="Hopper"
    )
    access = exchange_token(runtime, poll["id"], token)
    submit_response(runtime, access, {poll["questions"][0]["id"]: "yes"})
    raw = (runtime / "polls" / poll["id"] / "invitations.json").read_text(
        encoding="utf-8"
    )
    rows = response_view(runtime, poll["id"])

    assert rows[0]["display_name"] == "Grace Hopper"
    assert "Grace" in raw
    assert "@" not in raw
    assert "email" not in raw.lower()


def test_concurrent_revisions_are_locked_and_server_assigned(tmp_path: Path) -> None:
    from mifp_app.services.polls import (
        exchange_token,
        submit_response,
    )

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    access = exchange_token(runtime, poll["id"], token)
    qid = poll["questions"][0]["id"]

    with ThreadPoolExecutor(max_workers=4) as executor:
        results = list(
            executor.map(
                lambda index: submit_response(
                    runtime, access, {qid: "yes" if index % 2 else "no"}
                ),
                range(8),
            )
        )

    assert sorted(row["revision"] for row in results) == list(range(1, 9))
    assert len({row["response_id"] for row in results}) == 8
    assert all(row["submitted_at"].endswith("+00:00") for row in results)


def test_revoked_wrong_expired_and_locked_invitations_are_rejected(tmp_path: Path) -> None:
    from mifp_app.services.polls import (
        PollError,
        create_invitation,
        exchange_token,
        invitation_rows,
        save_poll,
        set_invitation_status,
        submit_response,
    )

    runtime = tmp_path / "config"
    identified = open_poll(runtime)
    identified["response_mode"] = "identified"
    identified = save_poll(runtime, identified["id"], identified)
    invitation, token = create_invitation(runtime, identified["id"], first_name="Alan", last_name="Turing")
    set_invitation_status(runtime, identified["id"], invitation["invitation_id"], "revoked")
    revoked = invitation_rows(runtime, identified["id"])[0]
    assert revoked["token_hash"] == ""
    assert "display_name" not in revoked
    with pytest.raises(PollError):
        exchange_token(runtime, identified["id"], token)

    locked = open_poll(runtime)
    locked["allow_changes"] = False
    locked = save_poll(runtime, locked["id"], locked)
    _invitation, locked_token = create_sent_invitation(runtime, locked["id"])
    access = exchange_token(runtime, locked["id"], locked_token)
    qid = locked["questions"][0]["id"]
    submit_response(runtime, access, {qid: "yes"})
    with pytest.raises(PollError):
        submit_response(runtime, access, {qid: "no"})
    with pytest.raises(PollError):
        exchange_token(runtime, identified["id"], locked_token)

    expired = open_poll(runtime)
    expired["deadline"] = "2020-01-01"
    expired = save_poll(runtime, expired["id"], expired)
    with pytest.raises(PollError):
        create_invitation(runtime, expired["id"])


def test_pending_invitation_is_not_a_valid_bearer_token(tmp_path: Path) -> None:
    from mifp_app.services.polls import PollError, create_invitation, exchange_token

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    _invitation, token = create_invitation(runtime, poll["id"])
    with pytest.raises(PollError, match="invalid or no longer active"):
        exchange_token(runtime, poll["id"], token)


def test_respondent_session_duration_defaults_bounds_and_expiry(tmp_path: Path) -> None:
    from mifp_app.services.polls import (
        PollError,
        PollSessionExpired,
        exchange_token,
        get_poll,
        respondent_poll,
        save_poll,
    )

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    assert poll["respondent_session_minutes"] == 60
    poll["respondent_session_minutes"] = 30
    poll = save_poll(runtime, poll["id"], poll)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    access = exchange_token(runtime, poll["id"], token)
    issued = datetime.fromisoformat(access["issued_at"])
    expires = datetime.fromisoformat(access["expires_at"])
    assert access["session_minutes"] == 30
    assert expires - issued == timedelta(minutes=30)

    expired_access = {**access, "expires_at": "2000-01-01T00:00:00+00:00"}
    with pytest.raises(PollSessionExpired, match="session has expired"):
        respondent_poll(runtime, expired_access)

    invalid = dict(get_poll(runtime, poll["id"]))
    invalid["respondent_session_minutes"] = 9
    with pytest.raises(PollError):
        save_poll(runtime, poll["id"], invalid)
    invalid["respondent_session_minutes"] = 1441
    with pytest.raises(PollError):
        save_poll(runtime, poll["id"], invalid)


def test_response_history_and_daily_activity_survive_cleaning(
    tmp_path: Path, monkeypatch
) -> None:
    from mifp_app.services import polls

    runtime = tmp_path / "config"
    now = [datetime(2027, 3, 3, 10, 0, tzinfo=UTC)]
    monkeypatch.setattr(polls, "utc_now", lambda: now[0])
    poll = open_poll(runtime)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    qid = poll["questions"][0]["id"]

    access = polls.exchange_token(runtime, poll["id"], token)
    polls.submit_response(runtime, access, {qid: "yes"})
    now[0] = datetime(2027, 3, 5, 9, 30, tzinfo=UTC)
    access = polls.exchange_token(runtime, poll["id"], token)
    polls.submit_response(runtime, access, {qid: "no"})
    now[0] = datetime(2027, 3, 7, 17, 45, tzinfo=UTC)
    access = polls.exchange_token(runtime, poll["id"], token)
    polls.submit_response(runtime, access, {qid: "yes"})

    response = polls.response_view(runtime, poll["id"])[0]
    history = polls.response_history(runtime, poll["id"], response["response_id"])
    analysis = polls.poll_analysis(runtime, poll["id"])
    assert response["revision_count"] == 3
    assert response["first_submitted_at"].startswith("2027-03-03")
    assert response["last_modified_at"].startswith("2027-03-07")
    assert len(history["revisions"]) == 3
    assert history["revisions"][0]["changed_questions"] == ["Will you attend?"]
    assert analysis["modified_responses"] == 1
    assert analysis["total_revisions"] == 3
    assert analysis["activity"] == [
        {"date": "2027-03-03", "new": 1, "modified": 0},
        {"date": "2027-03-05", "new": 0, "modified": 1},
        {"date": "2027-03-07", "new": 0, "modified": 1},
    ]

    assert polls.clean_poll(runtime, poll["id"])["revisions_removed"] == 0
    assert len((runtime / "polls" / poll["id"] / "responses.jsonl").read_text().splitlines()) == 3
    preserved_response = polls.response_view(runtime, poll["id"])[0]
    preserved = polls.response_history(runtime, poll["id"], preserved_response["response_id"])
    preserved_analysis = polls.poll_analysis(runtime, poll["id"])
    assert preserved["history_compacted"] is False
    assert preserved["revision_count"] == 3
    assert len(preserved["revision_timestamps"]) == 3
    assert len(preserved["revisions"]) == 3
    assert preserved_analysis["activity"] == analysis["activity"]


def test_truncated_final_jsonl_is_tolerated_but_middle_corruption_is_not(
    tmp_path: Path,
) -> None:
    from mifp_app.services.polls import (
        PollError,
        current_responses,
        exchange_token,
        submit_response,
    )

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    access = exchange_token(runtime, poll["id"], token)
    submit_response(runtime, access, {poll["questions"][0]["id"]: "yes"})
    path = runtime / "polls" / poll["id"] / "responses.jsonl"
    with path.open("ab") as handle:
        handle.write(b'{"truncated":')
    assert len(current_responses(runtime, poll["id"])) == 1
    repaired = submit_response(runtime, access, {poll["questions"][0]["id"]: "no"})
    assert repaired["revision"] == 2
    assert len(current_responses(runtime, poll["id"])) == 1
    path.write_bytes(b"not-json\n" + path.read_bytes())
    with pytest.raises(PollError):
        current_responses(runtime, poll["id"])


def test_exports_use_current_revision_and_block_formula_injection(
    tmp_path: Path,
) -> None:
    from mifp_app.services.polls import (
        exchange_token,
        export_csv,
        export_pdf,
        export_xlsx,
        save_poll,
        submit_response,
    )

    runtime = tmp_path / "config"
    poll = open_poll(runtime, "blank")
    poll["questions"][0]["question"] = "Comment"
    poll = save_poll(runtime, poll["id"], poll)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    access = exchange_token(runtime, poll["id"], token)
    qid = poll["questions"][0]["id"]
    submit_response(runtime, access, {qid: "old"})
    submit_response(runtime, access, {qid: '=HYPERLINK("bad") <img src="file:///etc/passwd"> & private'})

    csv_payload = export_csv(runtime, poll["id"]).decode("utf-8-sig")
    assert "old" not in csv_payload
    assert "'=HYPERLINK" in csv_payload
    workbook = load_workbook(
        io.BytesIO(export_xlsx(runtime, poll["id"])), data_only=False
    )
    assert workbook.sheetnames == ["Overview", "Responses", "Analysis"]
    assert workbook["Responses"]["E2"].value.startswith("'=")
    pdf = export_pdf(runtime, poll["id"])
    assert pdf.startswith(b"%PDF-")
    assert b"Times-Bold" in pdf
    assert b"Helvetica" in pdf


def test_poll_dashboard_and_public_token_exchange_flow(poll_app, monkeypatch) -> None:
    from mifp_app.routes import dashboard_polls

    runtime = Path(poll_app.config["RUNTIME_CONFIG_DIR"])
    poll = open_poll(runtime)
    invitation, token = create_sent_invitation(runtime, poll["id"])
    dashboard_client = logged_in_client(poll_app)

    assert (
        poll_app.test_client().get("/dashboard/notifications/polls").status_code == 302
    )
    page = dashboard_client.get(f"/dashboard/notifications/polls/{poll['id']}")
    assert page.status_code == 200
    assert b"data-poll-editor" in page.data
    assert (
        b"Recipient file" not in page.data
    )  # Poll recipients use the explicit CSV/XLSX control.
    assert b"browser only" in page.data
    assert b"Single recipient" in page.data
    assert b"Send poll" in page.data
    assert b"Send test" not in page.data
    assert b"Respondent session" in page.data
    assert b"data-response-history" not in page.data  # no responses yet

    sent = []
    monkeypatch.setattr(
        dashboard_polls, "send_mail", lambda _app, **kwargs: sent.append(kwargs) or True
    )
    invite = dashboard_client.post(
        f"/dashboard/notifications/polls/{poll['id']}/invite",
        json={"email": "person@example.org", "first_name": "Person", "password": "secret123"},
    )
    assert invite.status_code == 200
    assert len(sent) == 1
    assert "#p=" in sent[0]["body"] and "&t=" in sent[0]["body"]
    assert sent[0]["privacy_safe_log"] is True
    html_body = sent[0]["html_body"]
    assert '<a href="' in html_body
    assert "#p=" in html_body and "&amp;t=" in html_body
    assert "display:inline-block" in html_body
    assert 'bgcolor="' in html_body
    assert "This message contains no open or click tracking." in html_body
    assert (
        "person@example.org"
        not in (runtime / "polls" / poll["id"] / "invitations.json").read_text()
    )

    public_client = poll_app.test_client()
    initial = public_client.get("/respond")
    assert initial.status_code == 200
    assert b'id="mainNav"' not in initial.data
    assert initial.headers["Cache-Control"].startswith("no-store")
    assert initial.headers["Referrer-Policy"] == "no-referrer"
    assert initial.headers["X-Robots-Tag"] == "noindex, nofollow"
    assert token.encode() not in initial.data
    exchanged = public_client.post(
        "/respond/exchange", json={"poll_id": poll["id"], "token": token}
    )
    assert exchanged.status_code == 200
    loaded = public_client.get("/respond/poll")
    assert loaded.status_code == 200
    assert "token_hash" not in loaded.get_data(as_text=True)
    submitted = public_client.post(
        "/respond/submit", json={"answers": {poll["questions"][0]["id"]: "yes"}}
    )
    assert submitted.status_code == 200
    exported = dashboard_client.get(
        f"/dashboard/notifications/polls/{poll['id']}/export.pdf"
    )
    assert exported.status_code == 200
    assert exported.mimetype == "application/pdf"
    assert exported.headers["Cache-Control"] == "no-store, max-age=0"
    assert exported.headers["X-Content-Type-Options"] == "nosniff"
    with public_client.session_transaction() as public_session:
        assert public_session.get("admin_logged_in") is None
        access = public_session["poll_access"]
        assert access["poll_id"] == poll["id"]
        assert access["invitation_id"] == invitation["invitation_id"]
        assert access["session_minutes"] == 60
        assert access["issued_at"]
        assert access["expires_at"]


def test_public_poll_exchange_uses_existing_csrf_bootstrap(poll_app) -> None:

    runtime = Path(poll_app.config["RUNTIME_CONFIG_DIR"])
    poll = open_poll(runtime)
    _invitation, token = create_sent_invitation(runtime, poll["id"])
    poll_app.config["WTF_CSRF_ENABLED"] = True
    client = poll_app.test_client()
    page = client.get("/respond")
    match = re.search(rb'data-csrf-token="([^"]+)"', page.data)
    assert match
    csrf = match.group(1).decode()

    rejected = client.post(
        "/respond/exchange", json={"poll_id": poll["id"], "token": token}
    )
    accepted = client.post(
        "/respond/exchange",
        json={"poll_id": poll["id"], "token": token},
        headers={"X-CSRF-Token": csrf},
    )
    assert rejected.status_code == 400
    assert accepted.status_code == 200


def test_ephemeral_email_endpoint_personalizes_safely_without_history(
    poll_app, monkeypatch
) -> None:
    from mifp_app.routes import dashboard_notifications

    delivered = []
    monkeypatch.setattr(
        dashboard_notifications,
        "send_mail",
        lambda _app, **kwargs: delivered.append(kwargs) or True,
    )
    client = logged_in_client(poll_app)
    response = client.post(
        "/dashboard/notifications/send-one",
        json={
            "password": "secret123",
            "email": "recipient@example.org",
            "first_name": '<img src="https://tracker.invalid/x">',
            "last_name": "=Formula",
            "subject": "Hello {{first_name}}",
            "mail_title": "MIFP update",
            "message": "Hello {{first_name}},\n\nUpdate for {{last_name}}.",
            "message_html": "<p>Hello {{first_name}}</p><p>{{last_name}}</p>",
        },
    )

    assert response.status_code == 200
    assert len(delivered) == 1
    assert '<img src="https://tracker.invalid/x">' not in delivered[0]["html_body"]
    assert "&lt;img" in delivered[0]["html_body"]
    assert delivered[0]["to"] == "recipient@example.org"
    assert delivered[0]["privacy_safe_log"] is True
    assert not (
        Path(poll_app.config["RUNTIME_CONFIG_DIR"]) / "manual_email_history.json"
    ).exists()
    log_text = "\n".join(
        path.read_text(encoding="utf-8", errors="replace")
        for path in Path(poll_app.config["LOG_DIR"]).glob("*")
        if path.is_file()
    )
    assert "recipient@example.org" not in log_text
    assert "tracker.invalid" not in log_text


def test_lifecycle_clean_anonymize_reset_delete_and_retention(
    tmp_path: Path, monkeypatch
) -> None:
    from mifp_app.services import polls

    runtime = tmp_path / "config"
    poll = open_poll(runtime)
    poll["response_mode"] = "identified"
    poll = polls.save_poll(runtime, poll["id"], poll)
    _invitation, token = create_sent_invitation(
        runtime, poll["id"], first_name="Marie", last_name="Curie"
    )
    access = polls.exchange_token(runtime, poll["id"], token)
    qid = poll["questions"][0]["id"]
    polls.submit_response(runtime, access, {qid: "yes"})
    polls.submit_response(runtime, access, {qid: "no"})
    assert polls.clean_poll(runtime, poll["id"])["revisions_removed"] == 0
    assert len((runtime / "polls" / poll["id"] / "responses.jsonl").read_text().splitlines()) == 2
    current = polls.response_view(runtime, poll["id"])[0]
    assert len(polls.response_history(runtime, poll["id"], current["response_id"])["revisions"]) == 2
    polls.anonymize_poll(runtime, poll["id"])
    assert polls.get_poll(runtime, poll["id"])["response_mode"] == "anonymous"
    assert (
        "Marie" not in (runtime / "polls" / poll["id"] / "invitations.json").read_text()
    )
    polls.reset_poll(runtime, poll["id"])
    assert polls.current_responses(runtime, poll["id"]) == []
    polls.delete_poll(runtime, poll["id"])
    assert not (runtime / "polls" / poll["id"]).exists()

    expired = open_poll(runtime)
    expired["deadline"] = "2019-12-01"
    expired["retention_until"] = "2020-01-01"
    expired["status"] = "closed"
    polls.save_poll(runtime, expired["id"], expired)
    assert polls.cleanup_retention(runtime) == 1
    cleaned = polls.get_poll(runtime, expired["id"])
    assert cleaned["retention_cleaned_at"]


def test_poll_ids_reject_path_traversal(tmp_path: Path) -> None:
    from mifp_app.services.polls import PollNotFound, get_poll

    with pytest.raises(PollNotFound):
        get_poll(tmp_path / "config", "../../outside")


def test_client_recipient_parser_never_uses_web_storage_or_upload_field(poll_app) -> None:
    root = Path(__file__).resolve().parents[2]
    parser = (
        root / "MIFPAPP/CORE/mifp_app/static/js/dashboard/recipient-files.js"
    ).read_text(encoding="utf-8")
    email_batch = (
        root / "MIFPAPP/CORE/mifp_app/static/js/dashboard/email-batch.js"
    ).read_text(encoding="utf-8")
    template = logged_in_client(poll_app).get("/dashboard/notifications/batch").get_data(as_text=True)
    poll_template = (
        root / "MIFPAPP/CORE/mifp_app/templates/dashboard/polls.html"
    ).read_text(encoding="utf-8")

    assert "localStorage" not in parser + email_batch
    assert "sessionStorage" not in parser + email_batch
    assert "indexedDB" not in (parser + email_batch).lower()
    assert "data-email-recipient-file" in template
    assert "Batch email" in template
    direct = logged_in_client(poll_app).get("/dashboard/notifications").get_data(as_text=True)
    assert "data-email-recipient-file" not in direct
    assert "data-recipient-template-download" in template
    assert "data-recipient-template-download" in poll_template
    assert "downloadTemplate" in parser
    assert "email,first_name,last_name" in parser
    recipient_input = re.search(
        r"<input[^>]+data-email-recipient-file[^>]*>", template
    ).group(0)
    assert "name=" not in recipient_input
    assert "Promise.all" not in email_batch

def test_poll_invitation_email_uses_external_mifp_layout_without_tracking() -> None:
    from mifp_app.services.notifications import render_poll_invitation_html

    rendered = render_poll_invitation_html(
        title="Conference availability",
        body="Hello, please choose your preferred date.",
        body_html=None,
        cta_url="https://mifp.eu/respond#p=abc&t=secret",
        cta_label="Open poll",
        deadline="2027-03-20",
        question_count=3,
        allow_changes=True,
    )

    assert "Mediterranean Institute of Fundamental Physics" in rendered
    assert "Poll invitation" in rendered
    assert "Conference availability" in rendered
    assert "20 March 2027" in rendered
    assert "3 questions" in rendered
    assert "Open poll" in rendered
    assert "no open or click tracking" in rendered
    assert "Sent from the MIFP administrative dashboard" not in rendered
    assert "<script" not in rendered.lower()
    assert "<img" not in rendered.lower()



@pytest.mark.parametrize("password", [None, "wrong-confirmation-secret", "secret123"])
@pytest.mark.parametrize("configured_hash", ["valid", "", "malformed"])
def test_poll_send_confirmation_fails_closed_before_creating_tokens(
    poll_app, monkeypatch, password, configured_hash
) -> None:
    from mifp_app.routes import dashboard_polls
    from mifp_app.services import polls

    runtime = Path(poll_app.config["RUNTIME_CONFIG_DIR"])
    poll = open_poll(runtime)
    before = {path.name: path.read_bytes() for path in (runtime / "polls" / poll["id"]).iterdir() if path.is_file()}
    if configured_hash != "valid":
        poll_app.config["ADMIN_PASSWORD_HASH"] = configured_hash
    delivered = []
    monkeypatch.setattr(dashboard_polls, "send_mail", lambda *_a, **kw: delivered.append(kw) or True)
    payload = {"email": "person@example.org"}
    if password is not None:
        payload["password"] = password
    response = logged_in_client(poll_app).post(f"/dashboard/notifications/polls/{poll['id']}/invite", json=payload)
    if password == "secret123" and configured_hash == "valid":
        assert response.status_code == 200
        assert len(delivered) == 1
        assert "password" not in delivered[0]
        assert "secret123" not in str(delivered[0])
        assert polls.get_poll(runtime, poll["id"])["id"] == poll["id"]
    else:
        assert response.status_code == 403
        assert delivered == []
        after = {path.name: path.read_bytes() for path in (runtime / "polls" / poll["id"]).iterdir() if path.is_file()}
        assert after == before
    assert "wrong-confirmation-secret" not in response.get_data(as_text=True)
    for key in ("LOG_DIR", "RUNTIME_CONFIG_DIR"):
        for path in Path(poll_app.config[key]).rglob("*"):
            if path.is_file():
                assert b"wrong-confirmation-secret" not in path.read_bytes()


def test_poll_invitation_confirmation_preserves_auth_and_csrf(poll_app, monkeypatch) -> None:
    from mifp_app.routes import dashboard_polls

    poll = open_poll(Path(poll_app.config["RUNTIME_CONFIG_DIR"]))
    url = f"/dashboard/notifications/polls/{poll['id']}/invite"
    delivered = []
    monkeypatch.setattr(dashboard_polls, "send_mail", lambda *_a, **kw: delivered.append(kw) or True)
    assert poll_app.test_client().post(url, json={"password": "secret123"}).status_code in {302, 401}
    poll_app.config["WTF_CSRF_ENABLED"] = True
    response = logged_in_client(poll_app).post(url, json={"email": "person@example.org", "password": "secret123"})
    assert response.status_code == 400
    assert delivered == []

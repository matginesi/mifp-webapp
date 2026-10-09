"""Synthetic registrations and mail stubs only; no host or participant mail."""
import csv
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image

from TESTS.webapp.test_registrations import (storage, registration_client, authenticate,
                                            EVENT, ID, FIELDS, host_module, host_args)  # noqa: F401
from mifp_app.services.registrations import read_records
from mifp_app.services.registration_deletion import delete_registrations, DeletionRecoveryRequired

BASE = '/dashboard/conference-sites/1'
PASSWORD = 'synthetic-operation-password'


def post_data(**values):
    return dict(_csrf_token='synthetic-csrf', password=PASSWORD, confirm=EVENT, **values)


def add_record(root, receipt, proof=True):
    path = root / 'registrations' / EVENT / 'registrations.csv'
    original = read_records(root, EVENT)[0]
    row = dict(original, receipt_id=receipt, proof_file=f'proof-{receipt}.png' if proof else '')
    with path.open('a', newline='') as stream:
        csv.DictWriter(stream, fieldnames=FIELDS).writerow(row)
    if proof:
        Image.new('RGB', (100, 150), 'green').save(root / 'uploads' / EVENT / row['proof_file'])
    return row


@pytest.fixture
def deletion_client(registration_client, monkeypatch):
    from mifp_app.routes import dashboard_registrations as route
    app, client = registration_client
    app.config.update(REGISTRATIONS_WRITE_ENABLED=True, MAIL_PROVIDER='smtp')
    authenticate(client)
    alerts = []
    def alert(*args, **kwargs):
        alerts.append(kwargs)
        return SimpleNamespace(sent=True)
    monkeypatch.setattr(route, 'notify', alert)
    return app, client, alerts


@pytest.mark.parametrize('scope', ['single', 'multiple', 'all'])
def test_password_deletes_records_and_proofs_and_sends_alerts(deletion_client, storage, scope):
    app, client, alerts = deletion_client
    root, _ = storage
    second = add_record(root, 'REG-20261009-00000002')
    third = add_record(root, 'REG-20261009-00000003', proof=False)
    selected = [ID] if scope == 'single' else [ID, second['receipt_id']]
    expected = 3 if scope == 'all' else len(selected)
    response = client.post(BASE + '/registrations-delete?q=ignored-filter',
                           data=post_data(receipt=selected, all='1' if scope == 'all' else '0'))
    assert response.status_code == 302
    assert len(read_records(root, EVENT)) == 3 - expected
    assert not (root / 'uploads' / EVENT / f'proof-{ID}.pdf').exists()
    assert (root / 'uploads' / EVENT / second['proof_file']).exists() == (scope == 'single')
    assert [a['event'] for a in alerts] == ['registrations.deletion_authorized', 'registrations.deletion_completed']
    assert alerts[1]['metadata']['registration_count'] == expected
    assert all('synthetic@example.invalid' not in a['body'] and PASSWORD not in a['body'] for a in alerts)
    operation_root = Path(app.config['CONFERENCES_DIR']) / '.registrations-admin/deletion-operations'
    journals = list(operation_root.glob('*.json'))
    assert len(journals) == 2
    assert all(path.stat().st_mode & 0o777 == 0o600 for path in journals)
    assert all(set(json.loads(path.read_text())) <= {'event', 'count', 'operation', 'status'} for path in journals)
    assert not list(Path(app.config['CONFERENCES_DIR']).rglob('*.zip'))
    assert not list(Path(app.config['CONFERENCES_DIR']).rglob('registrations.csv'))



@pytest.mark.parametrize('changes,code', [({'password': 'wrong'},403), ({'confirm':'wrong'},400),
                                         ({'receipt':'unknown'},400), ({'receipt':[]},400)])
def test_delete_denied_without_changing_files_or_sending_alerts(deletion_client, storage, changes, code):
    _, client, alerts = deletion_client
    root, row = storage
    data = post_data(receipt=ID); data.update(changes)
    assert client.post(BASE+'/registrations-delete', data=data).status_code == code
    assert read_records(root, EVENT) == [row]
    assert alerts == []


def test_deletion_requires_write_opt_in_and_working_alert(deletion_client, storage, monkeypatch):
    from mifp_app.routes import dashboard_registrations as route
    app, client, alerts = deletion_client
    root, row = storage
    app.config['REGISTRATIONS_WRITE_ENABLED'] = False
    assert client.post(BASE+'/registrations-delete',data=post_data(receipt=ID)).status_code == 409
    assert not alerts
    app.config['REGISTRATIONS_WRITE_ENABLED'] = True
    monkeypatch.setattr(route, 'notify', lambda *a, **k: SimpleNamespace(sent=False))
    response = client.post(BASE+'/registrations-delete',data=post_data(receipt=ID))
    assert response.status_code == 503 and b'No data was deleted' in response.data
    assert read_records(root, EVENT) == [row]


def test_deletion_requires_csrf_and_rate_limits(deletion_client, storage, monkeypatch):
    from mifp_app.routes import dashboard_registrations as route
    _, client, alerts = deletion_client
    assert client.post(BASE+'/registrations-delete',data={'password':PASSWORD,'receipt':ID,'confirm':EVENT}).status_code == 400
    authenticate(client)
    monkeypatch.setattr(route, 'ip_rate_allowed',lambda *a,**k:False)
    assert client.post(BASE+'/registrations-delete',data=post_data(receipt=ID)).status_code == 429
    assert not alerts


def test_deletion_journal_failure_is_non_destructive(storage, tmp_path, monkeypatch):
    from mifp_app.services import registration_deletion as service
    root, row = storage
    monkeypatch.setattr(service, '_journal', lambda *a, **k: (_ for _ in ()).throw(OSError('disk full')))
    with pytest.raises(OSError):
        delete_registrations(root, EVENT, [ID], tmp_path/'operations')
    assert read_records(root, EVENT) == [row]
    assert (root/'uploads'/EVENT/row['proof_file']).exists()


def test_unlink_failure_restores_csv_and_already_removed_files(storage, tmp_path, monkeypatch):
    from mifp_app.services import registration_deletion as service
    root, row = storage
    (root/'registrations'/EVENT/(ID+'.notifications.json')).write_text('{}')
    original_unlink = os.unlink
    def fail(name, *args, **kwargs):
        if str(name).endswith('.notifications.json'):
            raise OSError('synthetic unlink failure')
        return original_unlink(name,*args,**kwargs)
    monkeypatch.setattr(service.os, 'unlink', fail)
    with pytest.raises(OSError):
        delete_registrations(root, EVENT, [ID], tmp_path/'operations')
    assert read_records(root, EVENT) == [row]
    assert (root/'uploads'/EVENT/row['proof_file']).read_bytes() == b'%PDF-synthetic'
    assert json.loads(next((tmp_path/'operations').glob('*.finished.json')).read_text())['status']=='rolled_back'


def test_unfinished_journal_blocks_new_deletion(storage, tmp_path):
    root, row = storage
    operations=tmp_path/'operations';operations.mkdir(mode=0o700)
    (operations/'old.started.json').write_text('{}')
    with pytest.raises(DeletionRecoveryRequired):
        delete_registrations(root, EVENT, [ID], operations)
    assert read_records(root, EVENT) == [row]


@pytest.mark.parametrize('target', ['proof','directory','hardlink'])
def test_delete_rejects_symlinks_and_hardlinks(storage,tmp_path,target):
    root,row=storage
    proof=root/'uploads'/EVENT/row['proof_file']
    if target=='directory':
        path=root/'uploads'/EVENT;path.rename(path.with_name('real'));path.symlink_to(path.with_name('real'))
    elif target=='proof':
        proof.unlink();proof.symlink_to('/etc/passwd')
    else:
        os.link(proof, tmp_path/'other-proof')
    with pytest.raises((OSError, ValueError)):
        delete_registrations(root, EVENT, [ID], tmp_path/'operations')
    assert read_records(root, EVENT)==[row]


@pytest.mark.parametrize('fmt', ['csv','xlsx','docx','pdf'])
def test_all_exports_require_password_post_and_csrf(registration_client,fmt):
    _,client=registration_client;authenticate(client)
    endpoint=BASE+'/registrations-export/'+fmt
    assert client.get(endpoint).status_code==405
    assert client.post(endpoint,data=post_data() | {'password':'wrong'}).status_code==403
    assert client.post(endpoint,data={'password':PASSWORD}).status_code==400


def test_document_exports_append_all_pdf_and_image_pages(registration_client,storage):
    import pypdfium2 as pdfium
    from reportlab.pdfgen import canvas
    from docx import Document
    from openpyxl import load_workbook
    _,client=registration_client;authenticate(client)
    root,row=storage
    pdf=io.BytesIO();writer=canvas.Canvas(pdf)
    for label in ('synthetic receipt page one', 'synthetic receipt page two'):
        writer.drawString(40,700,label);writer.showPage()
    writer.save();(root/'uploads'/EVENT/row['proof_file']).write_bytes(pdf.getvalue())
    second=add_record(root,'REG-20261009-00000002')
    for fmt in ('docx','pdf','xlsx','csv'):
        response=client.post(BASE+'/registrations-export/'+fmt,data=post_data())
        assert response.status_code==200, response.data[:300]
        assert 'no-store' in response.headers['Cache-Control']
        if fmt=='docx':
            doc=Document(io.BytesIO(response.data))
            text='\n'.join(p.text for p in doc.paragraphs)
            assert ID in text and second['receipt_id'] in text and 'Page 2 of 2' in text
            assert len(doc.inline_shapes)==3
            assert doc.tables[0].cell(0,0)._element.xml.count('B42318')>0
        elif fmt=='pdf':
            with pdfium.PdfDocument(response.data) as doc:
                assert len(doc)==4  # list followed by every receipt page
                page=doc[len(doc)-1];textpage=page.get_textpage()
                assert second['receipt_id'] in textpage.get_text_range()
                textpage.close();page.close()
        elif fmt=='xlsx':
            workbook=load_workbook(io.BytesIO(response.data));assert len(workbook.active._images)==0
            assert workbook.active['A4'].fill.fgColor.rgb.endswith('B42318')
        else:
            assert b"'=1+1" in response.data


def test_document_filter_and_missing_proof_notice(registration_client,storage):
    from docx import Document
    _,client=registration_client;authenticate(client)
    root,row=storage
    (root/'uploads'/EVENT/row['proof_file']).unlink()
    response=client.post(BASE+'/registrations-export/docx',data=post_data())
    doc=Document(io.BytesIO(response.data))
    assert any('missing' in p.text for p in doc.paragraphs)
    assert ID in '\n'.join(p.text for p in doc.paragraphs)
    response=client.post(BASE+'/registrations-export/docx?q=no-matching-record',data=post_data())
    assert response.status_code==200 and len(Document(io.BytesIO(response.data)).inline_shapes)==0


def test_host_write_access_requires_updated_runtime_and_revocation(storage,tmp_path):
    root,row=storage;host=host_module();args=host_args(tmp_path,root,row)
    args.request=EVENT;args.php_uid=os.getuid();args.php_gid=os.getgid();args.writable=True
    args.public_root=tmp_path/'published'
    source=args.public_root/EVENT/'regform/src';source.mkdir(parents=True)
    (source/'bootstrap.php').write_text('outdated runtime')
    (source.parent/'index.php').write_text('$registrationOperationLock .operations.lock')
    with pytest.raises(ValueError,match='Update'):
        host.grant_read(args)
    (source/'bootstrap.php').write_text('function registration_file_mode()')
    args.allowlist.write_text(EVENT+'/regform\n')
    with pytest.raises(ValueError,match='Revoke'):
        host.grant_read(args)
    args.allowlist.write_text('');host.grant_read(args)
    assert (root/'registrations'/EVENT).stat().st_mode & 0o7777==0o2770
    assert (root/'registrations'/EVENT/'registrations.csv').stat().st_mode & 0o777==0o660
    assert (root/'registrations'/EVENT/'.operations.lock').stat().st_mode & 0o777==0o660


def test_busy_php_lock_prevents_mutation(storage,tmp_path):
    import fcntl
    root,row=storage
    with (root/'registrations'/EVENT/'.operations.lock').open('rb') as stream:
        fcntl.flock(stream,fcntl.LOCK_EX)
        with pytest.raises(BlockingIOError):
            delete_registrations(root,EVENT,[ID],tmp_path/'operations')
    assert read_records(root,EVENT)==[row]
    assert not list((tmp_path/'operations').glob('*.zip'))


def test_partial_csv_write_is_rolled_back(storage,tmp_path,monkeypatch):
    from mifp_app.services import registration_deletion as service
    root,row=storage
    original=service._write_csv
    calls=[]
    def partial(stream,payload):
        calls.append(1)
        if len(calls)==1:
            stream.seek(0);stream.write(b'broken');stream.flush()
            raise OSError('synthetic write error')
        return original(stream,payload)
    monkeypatch.setattr(service,'_write_csv',partial)
    with pytest.raises(OSError):
        delete_registrations(root,EVENT,[ID],tmp_path/'operations')
    assert read_records(root,EVENT)==[row]
    assert (root/'uploads'/EVENT/row['proof_file']).exists()


def test_completion_alert_failure_does_not_repeat_deletion(deletion_client,storage,monkeypatch):
    from mifp_app.routes import dashboard_registrations as route
    _,client,_=deletion_client
    calls=[]
    def alert(*args,**kwargs):
        calls.append(kwargs['event']);return SimpleNamespace(sent=len(calls)==1)
    monkeypatch.setattr(route,'notify',alert)
    response=client.post(BASE+'/registrations-delete',data=post_data(receipt=ID),follow_redirects=True)
    assert response.status_code==200 and b'completion alert failed' in response.data
    assert read_records(storage[0],EVENT)==[]
    assert len(calls)==2


def test_document_size_limit_refuses_without_silently_omitting_receipts(registration_client,monkeypatch):
    from mifp_app.services import registration_exports
    _,client=registration_client;authenticate(client)
    monkeypatch.setattr(registration_exports,'MAX_EXPORT_BYTES',1)
    response=client.post(BASE+'/registrations-export/pdf',data=post_data())
    assert response.status_code==413
    assert client.post(BASE+'/registrations-export/csv',data=post_data()).status_code==200


def test_export_formula_prefix_and_markup_remain_inert(registration_client,storage):
    from openpyxl import load_workbook
    _,client=registration_client;authenticate(client)
    root,_=storage
    rows=read_records(root,EVENT);rows[0]['first_name']=' =HYPERLINK("https://example.invalid")'
    with (root/'registrations'/EVENT/'registrations.csv').open('w',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=FIELDS);writer.writeheader();writer.writerows(rows)
    response=client.post(BASE+'/registrations-export/xlsx',data=post_data())
    workbook=load_workbook(io.BytesIO(response.data))
    assert workbook.active.cell(5,3).data_type!='f'


def test_console_transport_cannot_authorize_real_deletion(deletion_client,storage):
    app,client,alerts=deletion_client
    app.config['MAIL_PROVIDER']='console'
    assert client.post(BASE+'/registrations-delete',data=post_data(receipt=ID)).status_code==503
    assert read_records(storage[0],EVENT)==[storage[1]] and not alerts


def test_deletion_reuses_real_notification_pipeline_with_stubbed_smtp(registration_client,storage,monkeypatch):
    from mifp_app.services import notifications
    from mifp_app.db.connection import connect
    app,client=registration_client;authenticate(client)
    app.config.update(REGISTRATIONS_WRITE_ENABLED=True,MAIL_PROVIDER='smtp',
                      SMTP_HOST='smtp.example.invalid',MAIL_FROM='noreply@example.invalid',
                      MAIL_TO='operator@example.invalid',SMTP_USERNAME='',SMTP_PASSWORD='')
    with connect(app.config['DATABASE_PATH']) as conn:
        conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('notification_enabled','0')")
        conn.commit()
    messages=[]
    def mail(*args,**kwargs):
        messages.append(kwargs);return True
    monkeypatch.setattr(notifications,'send_mail',mail)
    assert client.post(BASE+'/registrations-delete',data=post_data(receipt=ID)).status_code==302
    assert len(messages)==2 and all(m['to']=='operator@example.invalid' for m in messages)
    assert all('html_body' in m and PASSWORD not in m['body'] for m in messages)
    assert read_records(storage[0],EVENT)==[]


def test_deletion_preserves_preexisting_recovery_files(deletion_client, storage):
    app, client, _ = deletion_client
    old_directory = Path(app.config['CONFERENCES_DIR']) / '.registrations-admin/deletion-backups'
    old_directory.mkdir(parents=True, mode=0o700, exist_ok=True)
    old_copy = old_directory / 'preexisting.zip'
    old_copy.write_bytes(b'synthetic existing archive')
    assert client.post(BASE+'/registrations-delete', data=post_data(receipt=ID)).status_code == 302
    assert old_copy.read_bytes() == b'synthetic existing archive'
    assert list(old_directory.iterdir()) == [old_copy]


def test_failed_in_memory_rollback_blocks_next_deletion(storage, tmp_path, monkeypatch):
    from mifp_app.services import registration_deletion as service
    root, row = storage
    (root/'registrations'/EVENT/(ID+'.notifications.json')).write_text('{}')
    original_unlink, original_open = os.unlink, os.open
    def unlink(name, *args, **kwargs):
        if str(name).endswith('.notifications.json'):
            raise OSError('synthetic unlink failure')
        return original_unlink(name, *args, **kwargs)
    def opened(name, flags, *args, **kwargs):
        if str(name) == row['proof_file'] and flags & os.O_CREAT:
            raise OSError('synthetic rollback failure')
        return original_open(name, flags, *args, **kwargs)
    monkeypatch.setattr(service.os, 'unlink', unlink)
    monkeypatch.setattr(service.os, 'open', opened)
    with pytest.raises(DeletionRecoveryRequired):
        delete_registrations(root, EVENT, [ID], tmp_path/'operations')
    assert len(list((tmp_path/'operations').glob('*.started.json'))) == 1
    assert not list((tmp_path/'operations').glob('*.finished.json'))
    assert not list((tmp_path/'operations').glob('*.zip'))
    with pytest.raises(DeletionRecoveryRequired):
        delete_registrations(root, EVENT, [ID], tmp_path/'operations')
